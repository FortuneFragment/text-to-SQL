from __future__ import annotations

import csv
from html.parser import HTMLParser
import io
import json
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy.orm import Session

from repositories.text2sql_schema_annotation_repo import Text2SQLSchemaAnnotationRepository
from services.text2sql.config_service import Text2SQLConfigService


TABLE_HEADERS = {"table", "table_name", "table name", "tablename", "表名", "数据表", "物理表名"}
COLUMN_HEADERS = {"column", "column_name", "column name", "columnname", "name", "字段", "字段名", "列名", "物理列名"}
COLUMN_COMMENT_HEADERS = {
    "comment",
    "column_comment",
    "column comment",
    "columncomment",
    "字段注释",
    "字段说明",
    "字段含义",
    "中文名",
    "中文含义",
    "逻辑列名",
    "逻辑列名称",
}
TABLE_COMMENT_HEADERS = {
    "table_comment",
    "table comment",
    "tablecomment",
    "表注释",
    "表说明",
    "表含义",
    "逻辑表名",
    "逻辑表名称",
}
ALIASES_HEADERS = {"aliases", "alias", "同义词", "别名"}


class _HTMLTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.rows: list[list[str]] = []
        self._current_row: list[str] | None = None
        self._current_cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:  # noqa: ARG002
        lower = tag.lower()
        if lower == "tr":
            self._current_row = []
        elif lower in {"td", "th"} and self._current_row is not None:
            self._current_cell = []

    def handle_data(self, data: str) -> None:
        if self._current_cell is not None:
            self._current_cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        lower = tag.lower()
        if lower in {"td", "th"} and self._current_row is not None and self._current_cell is not None:
            self._current_row.append("".join(self._current_cell).strip())
            self._current_cell = None
        elif lower == "tr" and self._current_row is not None:
            if any(cell.strip() for cell in self._current_row):
                self.rows.append(self._current_row)
            self._current_row = None


class Text2SQLSchemaAnnotationImportService:
    def __init__(self, config_service: Text2SQLConfigService):
        self.config_service = config_service

    @staticmethod
    def _clean(value: Any) -> str:
        return str(value or "").strip()

    @classmethod
    def _header_key(cls, value: Any) -> str:
        return cls._clean(value).lower().replace("_", "").replace(" ", "")

    @classmethod
    def _pick(cls, row: dict[str, Any], candidates: set[str]) -> str:
        normalized = {cls._header_key(key): value for key, value in row.items()}
        for candidate in candidates:
            value = normalized.get(cls._header_key(candidate))
            if cls._clean(value):
                return cls._clean(value)
        return ""

    @classmethod
    def _split_aliases(cls, value: Any) -> list[str]:
        text = cls._clean(value)
        if not text:
            return []
        if text.startswith("["):
            try:
                parsed = json.loads(text)
            except json.JSONDecodeError:
                parsed = None
            if isinstance(parsed, list):
                return [cls._clean(item) for item in parsed if cls._clean(item)]
        separators = [";", "；", "，", "、", "|"]
        for separator in separators:
            text = text.replace(separator, ",")
        return [item.strip() for item in text.split(",") if item.strip()]

    @classmethod
    def _rows_from_table(cls, headers: list[Any], rows: list[list[Any]]) -> list[dict[str, Any]]:
        clean_headers = [cls._clean(header) for header in headers]
        result: list[dict[str, Any]] = []
        for raw_row in rows:
            if not any(cls._clean(cell) for cell in raw_row):
                continue
            result.append(
                {
                    clean_headers[index]: raw_row[index] if index < len(raw_row) else ""
                    for index in range(len(clean_headers))
                    if clean_headers[index]
                }
            )
        return result

    @classmethod
    def _load_xlsx(cls, content: bytes) -> list[dict[str, Any]]:
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except (BadZipFile, InvalidFileException):
            raise ValueError("上传的 Excel 不是有效的 .xlsx/.xlsm 文件，请重新另存为 .xlsx 或 CSV 后再上传") from None
        try:
            worksheet = workbook.worksheets[0]
            rows_iter = worksheet.iter_rows(values_only=True)
            try:
                headers = list(next(rows_iter))
            except StopIteration:
                return []
            rows = [list(row) for row in rows_iter]
            return cls._rows_from_table(headers, rows)
        finally:
            workbook.close()

    @classmethod
    def _load_xls(cls, content: bytes) -> list[dict[str, Any]]:
        try:
            import xlrd
        except ImportError:
            raise ValueError("当前文件是旧版 .xls，请另存为 .xlsx 或 CSV 后再上传") from None

        try:
            workbook = xlrd.open_workbook(file_contents=content)
            worksheet = workbook.sheet_by_index(0)
        except Exception:
            raise ValueError("上传的 .xls 文件无法解析，请另存为 .xlsx 或 CSV 后再上传") from None
        if worksheet.nrows <= 0:
            return []
        headers = worksheet.row_values(0)
        rows = [worksheet.row_values(index) for index in range(1, worksheet.nrows)]
        return cls._rows_from_table(headers, rows)

    @staticmethod
    def _decode_text(content: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk", "cp936", "big5"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise ValueError("文件编码无法识别，请另存为 UTF-8 CSV 后再上传")

    @staticmethod
    def _sniff_delimiter(text: str, fallback: str = ",") -> str:
        sample = "\n".join(line for line in text.splitlines()[:20] if line.strip())
        if not sample:
            return fallback
        try:
            return csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
        except csv.Error:
            candidates = [",", "\t", ";", "|"]
            return max(candidates, key=sample.count) if any(sample.count(item) for item in candidates) else fallback

    @classmethod
    def _load_delimited(cls, content: bytes, *, delimiter: str | None = None) -> list[dict[str, Any]]:
        text = cls._decode_text(content)
        if not text.strip():
            return []
        return [dict(row) for row in csv.DictReader(io.StringIO(text), delimiter=delimiter or cls._sniff_delimiter(text))]

    @classmethod
    def _load_html_table(cls, content: bytes) -> list[dict[str, Any]]:
        text = cls._decode_text(content)
        parser = _HTMLTableParser()
        parser.feed(text)
        if not parser.rows:
            return []
        return cls._rows_from_table(parser.rows[0], parser.rows[1:])

    @classmethod
    def _looks_like_html(cls, content: bytes) -> bool:
        try:
            prefix = cls._decode_text(content[:4096]).lstrip().lower()
        except ValueError:
            return False
        return prefix.startswith(("<!doctype html", "<html")) or "<table" in prefix

    @classmethod
    def _load_json(cls, content: bytes) -> list[dict[str, Any]]:
        payload = json.loads(cls._decode_text(content))
        if isinstance(payload, dict):
            payload = payload.get("annotations") or payload.get("items") or []
        if not isinstance(payload, list):
            raise ValueError("JSON must be a list or contain annotations/items list")
        return [dict(item) for item in payload if isinstance(item, dict)]

    @classmethod
    def _load_excel_like(cls, filename: str, content: bytes) -> list[dict[str, Any]]:
        if content.startswith(b"PK"):
            return cls._load_xlsx(content)
        if content.startswith(b"\xd0\xcf\x11\xe0"):
            return cls._load_xls(content)
        if cls._looks_like_html(content):
            return cls._load_html_table(content)
        try:
            rows = cls._load_delimited(content)
        except ValueError:
            suffix = ".xls" if str(filename or "").lower().endswith(".xls") else ".xlsx/.xlsm"
            raise ValueError(f"上传的 Excel 不是有效的 {suffix} 文件，请另存为 .xlsx 或 CSV 后再上传") from None
        if rows:
            return rows
        suffix = ".xls" if str(filename or "").lower().endswith(".xls") else ".xlsx/.xlsm"
        raise ValueError(f"上传的 Excel 不是有效的 {suffix} 文件，请另存为 .xlsx 或 CSV 后再上传")

    @classmethod
    def _load_rows(cls, filename: str, content: bytes) -> list[dict[str, Any]]:
        lower = str(filename or "").lower()
        if lower.endswith((".xlsx", ".xlsm", ".xls")):
            return cls._load_excel_like(filename, content)
        if lower.endswith(".csv"):
            return cls._load_delimited(content)
        if lower.endswith(".tsv"):
            return cls._load_delimited(content, delimiter="\t")
        if lower.endswith(".json"):
            return cls._load_json(content)
        if content.startswith(b"PK") or content.startswith(b"\xd0\xcf\x11\xe0") or cls._looks_like_html(content):
            return cls._load_excel_like(filename, content)
        raise ValueError("仅支持 .xlsx/.xlsm/.xls/.csv/.tsv/.json 字段注释文件")

    def import_file(
        self,
        db: Session,
        *,
        filename: str,
        content: bytes,
        connection_key: str | None = None,
    ) -> dict[str, Any]:
        resolved_key = self._clean(connection_key) or self.config_service.get_connection_key(db)
        if not resolved_key or resolved_key == "unconfigured":
            raise ValueError("请先配置并保存 SQL Server 业务库连接")
        if not content:
            raise ValueError("上传文件为空")

        raw_rows = self._load_rows(filename, content)
        items: list[dict[str, Any]] = []
        skipped = 0
        for row in raw_rows:
            table_name = self._pick(row, TABLE_HEADERS)
            column_name = self._pick(row, COLUMN_HEADERS)
            column_comment = self._pick(row, COLUMN_COMMENT_HEADERS)
            table_comment = self._pick(row, TABLE_COMMENT_HEADERS)
            aliases = self._split_aliases(self._pick(row, ALIASES_HEADERS))

            if not table_name or not column_name:
                skipped += 1
                continue

            item: dict[str, Any] = {
                "connection_key": resolved_key,
                "table_name": table_name,
                "column_name": column_name,
            }
            if column_comment:
                item["column_comment"] = column_comment
            if table_comment:
                item["table_comment"] = table_comment
            if aliases:
                item["aliases"] = aliases
            items.append(item)

        if not items:
            raise ValueError("未解析到有效字段注释，请检查表头是否包含 table_name、column_name、column_comment")

        Text2SQLSchemaAnnotationRepository(db).upsert_many(items)
        db.commit()
        return {
            "connection_key": resolved_key,
            "parsed": len(raw_rows),
            "upserted": len(items),
            "skipped": skipped,
            "sample": items[:10],
        }
