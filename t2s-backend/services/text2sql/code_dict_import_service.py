"""Text2SQL 码值字典轻量导入服务。

当前码值能力只需要两份最小结构：
  - 码值字典：category_key, code, name
  - 字段绑定：table_name, column_name, category_key

运行时不会基于这些表生成 JOIN；它们只用于生成 SQL 前注入 code->中文枚举，
以及 SQL 查询后把结果中的 code 替换成中文展示。
"""
from __future__ import annotations

import io
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from repositories.text2sql_code_dict_binding_repo import Text2SQLCodeDictBindingRepository
from repositories.text2sql_code_dict_value_repo import Text2SQLCodeDictValueRepository
from services.text2sql.config_service import Text2SQLConfigService


class Text2SQLCodeDictImportService:
    VALUE_CATEGORY_HEADERS = ("CATEGORY_KEY", "CATEGORY", "DICT_KEY", "字典编码", "字典", "类目编码", "类目")
    VALUE_CODE_HEADERS = ("CODE", "DM", "编码", "码值")
    VALUE_NAME_HEADERS = ("NAME", "MC", "名称", "中文", "中文名", "枚举值", "释义")

    BINDING_TABLE_HEADERS = ("TABLE_NAME", "TABLE", "源表", "业务表", "表名")
    BINDING_COLUMN_HEADERS = ("COLUMN_NAME", "COLUMN", "源列", "业务字段", "字段", "列名")
    BINDING_CATEGORY_HEADERS = VALUE_CATEGORY_HEADERS

    def __init__(self, config_service: Text2SQLConfigService):
        self.config_service = config_service

    @staticmethod
    def _clean(value: Any) -> str:
        return str(value or "").strip()

    @classmethod
    def _load_first_sheet(cls, content: bytes) -> tuple[list[str], list[list[Any]]]:
        """读 xlsx/xlsm 第一个工作表，返回 (表头, 数据行)。"""
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise ValueError(f"非法 Excel 文件（无法解析）：{exc}") from None
        try:
            worksheet = workbook.worksheets[0]
            rows_iter = worksheet.iter_rows(values_only=True)
            try:
                header = [cls._clean(cell) for cell in next(rows_iter)]
            except StopIteration:
                return [], []
            return header, [list(row) for row in rows_iter]
        finally:
            workbook.close()

    @staticmethod
    def _index_map(header: list[str]) -> dict[str, int]:
        return {str(name).strip().upper(): index for index, name in enumerate(header) if str(name).strip()}

    @staticmethod
    def _cell(row: list[Any], index: int) -> str:
        if index < 0 or index >= len(row) or row[index] is None:
            return ""
        return str(row[index]).strip()

    @classmethod
    def _pick(cls, indexes: dict[str, int], names: tuple[str, ...]) -> int:
        for name in names:
            index = indexes.get(str(name).strip().upper())
            if index is not None:
                return index
        return -1

    @classmethod
    def _classify_sheet(cls, header: list[str]) -> str:
        indexes = cls._index_map(header)
        has_binding = (
            cls._pick(indexes, cls.BINDING_TABLE_HEADERS) >= 0
            and cls._pick(indexes, cls.BINDING_COLUMN_HEADERS) >= 0
            and cls._pick(indexes, cls.BINDING_CATEGORY_HEADERS) >= 0
        )
        if has_binding:
            return "binding"

        has_value = (
            cls._pick(indexes, cls.VALUE_CATEGORY_HEADERS) >= 0
            and cls._pick(indexes, cls.VALUE_CODE_HEADERS) >= 0
            and cls._pick(indexes, cls.VALUE_NAME_HEADERS) >= 0
        )
        if has_value:
            return "value"
        return "unknown"

    @classmethod
    def build_value_items(cls, sheet: tuple[list[str], list[list[Any]]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        """从最小码值表构建取值项：category_key, code, name。"""
        header, rows = sheet
        indexes = cls._index_map(header)
        category_idx = cls._pick(indexes, cls.VALUE_CATEGORY_HEADERS)
        code_idx = cls._pick(indexes, cls.VALUE_CODE_HEADERS)
        name_idx = cls._pick(indexes, cls.VALUE_NAME_HEADERS)
        if min(category_idx, code_idx, name_idx) < 0:
            raise ValueError("码值字典缺少必要列：category_key/code/name")

        merged: dict[tuple[str, str], dict[str, Any]] = {}
        skipped = 0
        for row in rows:
            category_key = cls._clean(cls._cell(row, category_idx))
            code = cls._clean(cls._cell(row, code_idx))
            name = cls._clean(cls._cell(row, name_idx))
            if not category_key or not code or not name:
                skipped += 1
                continue
            merged[(category_key, code)] = {
                "category_key": category_key,
                "code": code,
                "name": name,
            }

        items = list(merged.values())
        return items, {
            "value_count": len(items),
            "category_count": len({item["category_key"] for item in items}),
            "skipped_rows": skipped,
        }

    @classmethod
    def build_binding_items(cls, sheet: tuple[list[str], list[list[Any]]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        """从最小字段绑定表构建绑定项：table_name, column_name, category_key。"""
        header, rows = sheet
        indexes = cls._index_map(header)
        table_idx = cls._pick(indexes, cls.BINDING_TABLE_HEADERS)
        column_idx = cls._pick(indexes, cls.BINDING_COLUMN_HEADERS)
        category_idx = cls._pick(indexes, cls.BINDING_CATEGORY_HEADERS)
        if min(table_idx, column_idx, category_idx) < 0:
            raise ValueError("字段绑定缺少必要列：table_name/column_name/category_key")

        merged: dict[tuple[str, str], dict[str, Any]] = {}
        skipped = 0
        for row in rows:
            table_name = cls._clean(cls._cell(row, table_idx)).upper()
            column_name = cls._clean(cls._cell(row, column_idx)).upper()
            category_key = cls._clean(cls._cell(row, category_idx))
            if not table_name or not column_name or not category_key:
                skipped += 1
                continue
            merged[(table_name, column_name)] = {
                "table_name": table_name,
                "column_name": column_name,
                "category_key": category_key,
            }

        items = list(merged.values())
        return items, {
            "binding_count": len(items),
            "skipped_rows": skipped,
        }

    def import_files(
        self,
        db: Session,
        *,
        files: list[dict[str, Any]],
        connection_key: str | None = None,
        expected_role: str | None = None,
    ) -> dict[str, Any]:
        """一次上传码值字典与字段绑定文件，分别落库到两张最小结构表。"""
        if expected_role and expected_role not in {"value", "binding"}:
            raise ValueError("不支持的码值字典导入类型")

        resolved_key = self._clean(connection_key) or self.config_service.get_connection_key(db)
        if not resolved_key or resolved_key == "unconfigured":
            raise ValueError("请先配置并保存 SQL Server 业务库连接")
        if not files:
            raise ValueError("未收到任何文件")

        value_items: list[dict[str, Any]] = []
        binding_items: list[dict[str, Any]] = []
        value_report = {"value_count": 0, "category_count": 0, "skipped_rows": 0}
        binding_report = {"binding_count": 0, "skipped_rows": 0}
        ignored_files: list[str] = []

        for item in files:
            filename = self._clean(item.get("filename")) or "未命名文件"
            content = item.get("content") or b""
            if not content:
                ignored_files.append(filename)
                continue
            if not filename.lower().endswith((".xlsx", ".xlsm")):
                ignored_files.append(filename)
                continue

            sheet = self._load_first_sheet(content)
            role = self._classify_sheet(sheet[0])
            if expected_role and role != expected_role:
                ignored_files.append(filename)
                continue

            if role == "value":
                items, report = self.build_value_items(sheet)
                value_items.extend(items)
                value_report["value_count"] += report["value_count"]
                value_report["skipped_rows"] += report["skipped_rows"]
            elif role == "binding":
                items, report = self.build_binding_items(sheet)
                binding_items.extend(items)
                binding_report["binding_count"] += report["binding_count"]
                binding_report["skipped_rows"] += report["skipped_rows"]
            else:
                ignored_files.append(filename)

        if not value_items and not binding_items:
            if expected_role == "value":
                raise ValueError("未识别到码值字典文件，请检查表头")
            if expected_role == "binding":
                raise ValueError("未识别到字段绑定文件，请检查表头")
            raise ValueError("未识别到码值字典或字段绑定文件，请检查表头")

        # 多个文件合并后按唯一键去重，后上传的行覆盖先上传的行。
        merged_values = {(item["category_key"], item["code"]): item for item in value_items}
        merged_bindings = {(item["table_name"], item["column_name"]): item for item in binding_items}
        value_items = list(merged_values.values())
        binding_items = list(merged_bindings.values())
        value_report["value_count"] = len(value_items)
        value_report["category_count"] = len({item["category_key"] for item in value_items})
        binding_report["binding_count"] = len(binding_items)

        for item in value_items:
            item["connection_key"] = resolved_key
        for item in binding_items:
            item["connection_key"] = resolved_key

        value_repo = Text2SQLCodeDictValueRepository(db)
        binding_repo = Text2SQLCodeDictBindingRepository(db)
        value_written = value_repo.bulk_replace(resolved_key, value_items) if value_items else 0
        binding_written = binding_repo.bulk_replace(resolved_key, binding_items) if binding_items else 0
        db.commit()

        return {
            "connection_key": resolved_key,
            "value": value_report | {"written": value_written},
            "binding": binding_report | {"written": binding_written},
            "ignored_files": ignored_files,
        }

    def import_value_files(
        self,
        db: Session,
        *,
        files: list[dict[str, Any]],
        connection_key: str | None = None,
    ) -> dict[str, Any]:
        """只导入码值字典文件，不影响字段绑定。"""
        return self.import_files(db, files=files, connection_key=connection_key, expected_role="value")

    def import_binding_files(
        self,
        db: Session,
        *,
        files: list[dict[str, Any]],
        connection_key: str | None = None,
    ) -> dict[str, Any]:
        """只导入字段绑定文件，不影响码值字典。"""
        return self.import_files(db, files=files, connection_key=connection_key, expected_role="binding")
