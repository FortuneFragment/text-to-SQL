from __future__ import annotations

import hashlib
import io
import json
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from repositories.text2sql_table_relation_repo import Text2SQLTableRelationRepository
from schemas.text2sql import (
    ColumnInfo,
    CreateText2SQLRelationRequest,
    Text2SQLRelationBatchImportResponse,
    Text2SQLRelationItem,
    Text2SQLRelationListResponse,
    Text2SQLRelationTableColumnsResponse,
    UpdateText2SQLRelationRequest,
)
from services.text2sql.config_service import GLOBAL_CONFIG_USER_ID, Text2SQLConfigService
from services.text2sql.schema_service import Text2SQLSchemaService


class Text2SQLRelationService:
    """Manage table join relations and expose relation hints for SQL generation."""

    XLSX_HEADERS = [
        "\u6e90\u8868",
        "\u6e90\u5217",
        "\u76ee\u6807\u8868",
        "\u76ee\u6807\u5217",
        "\u5173\u7cfb\u7c7b\u578b",
        "\u63cf\u8ff0",
    ]
    HEADER_ALIASES = {
        "\u6e90\u8868": "source_table",
        "\u4e3b\u8868": "source_table",
        "source_table": "source_table",
        "sourcetable": "source_table",
        "\u6e90\u5217": "source_columns",
        "\u6e90\u5b57\u6bb5": "source_columns",
        "\u4e3b\u8868\u5b57\u6bb5": "source_columns",
        "source_column": "source_columns",
        "source_columns": "source_columns",
        "sourcecolumns": "source_columns",
        "\u76ee\u6807\u8868": "target_table",
        "target_table": "target_table",
        "targettable": "target_table",
        "\u76ee\u6807\u5217": "target_columns",
        "\u76ee\u6807\u5b57\u6bb5": "target_columns",
        "target_column": "target_columns",
        "target_columns": "target_columns",
        "targetcolumns": "target_columns",
        "\u5173\u7cfb\u7c7b\u578b": "relation_type",
        "relation_type": "relation_type",
        "relationtype": "relation_type",
        "\u63cf\u8ff0": "description",
        "\u8bf4\u660e": "description",
        "\u5173\u7cfb\u63cf\u8ff0": "description",
        "description": "description",
    }

    def __init__(
        self,
        schema_service: Text2SQLSchemaService,
        config_service: Text2SQLConfigService,
    ):
        self.schema_service = schema_service
        self.config_service = config_service

    @staticmethod
    def _normalize_identifier(value: str | None) -> str:
        text = (value or "").strip().strip("`").strip('"').replace("[", "").replace("]", "")
        if "." in text:
            text = text.split(".")[-1]
        return text.lower()

    @staticmethod
    def _normalize_connection_key(value: str | None) -> str:
        text = str(value or "").strip()
        if len(text) <= 255:
            return text
        digest = hashlib.sha256(text.encode("utf-8")).hexdigest()
        return f"sha256:{digest}"

    @staticmethod
    def _cell_text(value: Any) -> str:
        if value is None:
            return ""
        return str(value).strip()

    @classmethod
    def _normalize_header(cls, value: Any) -> str:
        text = cls._cell_text(value)
        compact = text.replace(" ", "").replace("-", "_").replace(".", "_").lower()
        return cls.HEADER_ALIASES.get(text) or cls.HEADER_ALIASES.get(compact) or ""

    @classmethod
    def _split_column_text(cls, value: Any) -> list[str]:
        text = cls._cell_text(value)
        if not text:
            return []
        for separator in ("\uff0c", "\u3001", ",", ";", "\uff1b", "\n", "\r"):
            text = text.replace(separator, "|")
        return [part.strip() for part in text.split("|") if part.strip()]

    @staticmethod
    def _safe_load_columns(raw_value: str | None) -> list[str]:
        if not raw_value:
            return []
        try:
            parsed = json.loads(raw_value)
        except json.JSONDecodeError:
            return []
        if not isinstance(parsed, list):
            return []
        return [str(item).strip() for item in parsed if str(item).strip()]

    @classmethod
    def _normalize_columns(cls, columns: list[str]) -> list[str]:
        normalized_columns: list[str] = []
        seen: set[str] = set()
        for raw_column in columns:
            column_name = str(raw_column or "").strip()
            if not column_name:
                continue
            normalized = cls._normalize_identifier(column_name)
            if not normalized:
                continue
            if normalized in seen:
                raise ValueError(f"Duplicate column in composite key: {column_name}")
            seen.add(normalized)
            normalized_columns.append(column_name)
        return normalized_columns

    @classmethod
    def _build_columns_hash(cls, columns: list[str]) -> str:
        normalized = [cls._normalize_identifier(item) for item in columns if cls._normalize_identifier(item)]
        payload = "|".join(normalized)
        return hashlib.sha256(payload.encode("utf-8")).hexdigest()

    @classmethod
    def _format_relation_summary(
        cls,
        source_table: str,
        source_columns: list[str],
        target_table: str,
        target_columns: list[str],
    ) -> str:
        pairs = [
            f"{source_table}.{source_column} = {target_table}.{target_column}"
            for source_column, target_column in zip(source_columns, target_columns)
        ]
        return " AND ".join(pairs)

    def _resolve_table_name(self, db: Session, table_name: str) -> str:
        resolved_tables, missing_tables = self.schema_service.validate_selected_tables(db, [table_name])
        if missing_tables or not resolved_tables:
            raise ValueError(f"Table does not exist: {table_name}")
        return resolved_tables[0]

    def _resolve_columns(
        self,
        db: Session,
        table_name: str,
        columns: list[str],
    ) -> tuple[str, list[str]]:
        real_table_name = self._resolve_table_name(db, table_name)
        normalized_input_columns = self._normalize_columns(columns)
        if not normalized_input_columns:
            raise ValueError("Column list cannot be empty")

        table_detail = self.schema_service.get_table_detail(db, real_table_name)
        lookup = {
            self._normalize_identifier(str(column.get("name") or "")): str(column.get("name") or "")
            for column in table_detail.get("columns", [])
            if str(column.get("name") or "").strip()
        }
        resolved_columns: list[str] = []
        for raw_column in normalized_input_columns:
            normalized = self._normalize_identifier(raw_column)
            real_column = lookup.get(normalized)
            if not real_column:
                raise ValueError(f"Column does not exist: {real_table_name}.{raw_column}")
            resolved_columns.append(real_column)
        return real_table_name, resolved_columns

    @classmethod
    def _same_relation(
        cls,
        *,
        source_table: str,
        source_columns: list[str],
        target_table: str,
        target_columns: list[str],
        other_source_table: str,
        other_source_columns: list[str],
        other_target_table: str,
        other_target_columns: list[str],
    ) -> bool:
        left_source_table = cls._normalize_identifier(source_table)
        left_target_table = cls._normalize_identifier(target_table)
        left_source_columns = tuple(cls._normalize_identifier(item) for item in source_columns)
        left_target_columns = tuple(cls._normalize_identifier(item) for item in target_columns)
        right_source_table = cls._normalize_identifier(other_source_table)
        right_target_table = cls._normalize_identifier(other_target_table)
        right_source_columns = tuple(cls._normalize_identifier(item) for item in other_source_columns)
        right_target_columns = tuple(cls._normalize_identifier(item) for item in other_target_columns)
        return (
            left_source_table == right_source_table
            and left_target_table == right_target_table
            and left_source_columns == right_source_columns
            and left_target_columns == right_target_columns
        )

    def _ensure_no_duplicate_relation(
        self,
        db: Session,
        *,
        source_table: str,
        source_columns: list[str],
        target_table: str,
        target_columns: list[str],
        exclude_id: int | None = None,
    ) -> None:
        connection_key = self._normalize_connection_key(self.config_service.get_connection_key(db))
        repo = Text2SQLTableRelationRepository(db)
        existing = repo.list_by_pair(
            user_id=GLOBAL_CONFIG_USER_ID,
            connection_key=connection_key,
            source_table=source_table,
            target_table=target_table,
            exclude_id=exclude_id,
        )
        for item in existing:
            other_source_columns = self._safe_load_columns(item.source_columns)
            other_target_columns = self._safe_load_columns(item.target_columns)
            if self._same_relation(
                source_table=source_table,
                source_columns=source_columns,
                target_table=target_table,
                target_columns=target_columns,
                other_source_table=item.source_table,
                other_source_columns=other_source_columns,
                other_target_table=item.target_table,
                other_target_columns=other_target_columns,
            ):
                raise ValueError("Relation already exists")
            if self._same_relation(
                source_table=source_table,
                source_columns=source_columns,
                target_table=target_table,
                target_columns=target_columns,
                other_source_table=item.target_table,
                other_source_columns=other_target_columns,
                other_target_table=item.source_table,
                other_target_columns=other_source_columns,
            ):
                raise ValueError("Reverse relation already exists")

    def _build_relation_payload(
        self,
        db: Session,
        request: CreateText2SQLRelationRequest | UpdateText2SQLRelationRequest,
    ) -> dict:
        if len(request.source_columns) != len(request.target_columns):
            raise ValueError("source_columns and target_columns must have the same length")

        source_table, source_columns = self._resolve_columns(db, request.source_table, request.source_columns)
        target_table, target_columns = self._resolve_columns(db, request.target_table, request.target_columns)
        if len(source_columns) != len(target_columns):
            raise ValueError("Composite key column counts must match")

        relation_type = str(request.relation_type or "").strip() or "N:1"
        description = str(request.description or "").strip()
        connection_key = self._normalize_connection_key(self.config_service.get_connection_key(db))
        return {
            "user_id": GLOBAL_CONFIG_USER_ID,
            "connection_key": connection_key,
            "source_table": source_table,
            "source_columns": json.dumps(source_columns, ensure_ascii=False),
            "source_columns_hash": self._build_columns_hash(source_columns),
            "target_table": target_table,
            "target_columns": json.dumps(target_columns, ensure_ascii=False),
            "target_columns_hash": self._build_columns_hash(target_columns),
            "relation_type": relation_type,
            "description": description,
        }

    def _to_relation_item(self, relation) -> Text2SQLRelationItem:
        return Text2SQLRelationItem(
            id=int(relation.id),
            source_table=str(relation.source_table or ""),
            source_columns=self._safe_load_columns(relation.source_columns),
            target_table=str(relation.target_table or ""),
            target_columns=self._safe_load_columns(relation.target_columns),
            relation_type=str(relation.relation_type or ""),
            description=str(relation.description or ""),
            created_at=relation.created_at,
            updated_at=relation.updated_at,
        )

    def _get_owned_relation(self, db: Session, relation_id: int):
        repo = Text2SQLTableRelationRepository(db)
        relation = repo.get_by_id(relation_id)
        if relation is None:
            raise ValueError("Relation does not exist")
        connection_key = self._normalize_connection_key(self.config_service.get_connection_key(db))
        if int(relation.user_id) != GLOBAL_CONFIG_USER_ID or str(relation.connection_key) != connection_key:
            raise ValueError("Relation does not exist or is not accessible")
        return relation

    def list_relations(
        self,
        db: Session,
        *,
        page: int = 1,
        page_size: int = 20,
        keyword: str = "",
        table_name: str = "",
    ) -> Text2SQLRelationListResponse:
        connection_key = self._normalize_connection_key(self.config_service.get_connection_key(db))
        rows, total = Text2SQLTableRelationRepository(db).list_by_user_and_connection(
            user_id=GLOBAL_CONFIG_USER_ID,
            connection_key=connection_key,
            keyword=keyword,
            table_name=table_name,
            page=page,
            page_size=page_size,
        )
        return Text2SQLRelationListResponse(
            items=[self._to_relation_item(item) for item in rows],
            total=total,
            page=max(1, int(page)),
            page_size=max(1, int(page_size)),
        )

    def _parse_relations_xlsx(self, content: bytes) -> list[tuple[int, dict]]:
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise ValueError("Failed to parse XLSX file") from exc
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
        header_map: dict[str, int] = {}
        header_row_index = -1
        for row_index, row in enumerate(rows, start=1):
            mapped: dict[str, int] = {}
            for column_index, cell in enumerate(row):
                field_name = self._normalize_header(cell)
                if field_name and field_name not in mapped:
                    mapped[field_name] = column_index
            if {"source_table", "source_columns", "target_table", "target_columns"}.issubset(mapped):
                header_map = mapped
                header_row_index = row_index
                break
        if not header_map:
            raise ValueError("XLSX headers must include source/target table and column fields")

        parsed_rows: list[tuple[int, dict]] = []
        for row_index, row in enumerate(rows[header_row_index:], start=header_row_index + 1):
            if not any(self._cell_text(cell) for cell in row):
                continue

            def value(field_name: str) -> str:
                column_index = header_map.get(field_name)
                if column_index is None or column_index >= len(row):
                    return ""
                return self._cell_text(row[column_index])

            payload = {
                "source_table": value("source_table"),
                "source_columns": self._split_column_text(value("source_columns")),
                "target_table": value("target_table"),
                "target_columns": self._split_column_text(value("target_columns")),
                "relation_type": value("relation_type") or "N:1",
                "description": value("description"),
            }
            parsed_rows.append((row_index, payload))
        return parsed_rows

    def import_relations_from_xlsx(
        self,
        db: Session,
        *,
        filename: str,
        content: bytes,
    ) -> Text2SQLRelationBatchImportResponse:
        if not str(filename or "").lower().endswith((".xlsx", ".xlsm")):
            raise ValueError("Only .xlsx/.xlsm relation files are supported")
        parsed_rows = self._parse_relations_xlsx(content)
        summary = Text2SQLRelationBatchImportResponse(total=len(parsed_rows))
        for row_index, payload in parsed_rows:
            try:
                request = CreateText2SQLRelationRequest(**payload)
                self.create_relation(db, request)
                summary.created += 1
            except Exception as exc:  # noqa: BLE001
                summary.failed += 1
                if len(summary.errors) < 50:
                    summary.errors.append(f"Row {row_index}: {exc}")
        return summary

    def export_relations_xlsx(self, db: Session) -> bytes:
        connection_key = self._normalize_connection_key(self.config_service.get_connection_key(db))
        rows, _ = Text2SQLTableRelationRepository(db).list_by_user_and_connection(
            user_id=GLOBAL_CONFIG_USER_ID,
            connection_key=connection_key,
            page=1,
            page_size=100000,
        )
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "relations"
        worksheet.append(self.XLSX_HEADERS)
        for item in rows:
            worksheet.append(
                [
                    str(item.source_table or ""),
                    "|".join(self._safe_load_columns(item.source_columns)),
                    str(item.target_table or ""),
                    "|".join(self._safe_load_columns(item.target_columns)),
                    str(item.relation_type or ""),
                    str(item.description or ""),
                ]
            )
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def create_relation(self, db: Session, request: CreateText2SQLRelationRequest) -> Text2SQLRelationItem:
        payload = self._build_relation_payload(db, request)
        source_columns = self._safe_load_columns(payload["source_columns"])
        target_columns = self._safe_load_columns(payload["target_columns"])
        self._ensure_no_duplicate_relation(
            db,
            source_table=payload["source_table"],
            source_columns=source_columns,
            target_table=payload["target_table"],
            target_columns=target_columns,
        )
        relation = Text2SQLTableRelationRepository(db).create(payload)
        return self._to_relation_item(relation)

    def update_relation(
        self,
        db: Session,
        relation_id: int,
        request: UpdateText2SQLRelationRequest,
    ) -> Text2SQLRelationItem:
        relation = self._get_owned_relation(db, relation_id)
        payload = self._build_relation_payload(db, request)
        source_columns = self._safe_load_columns(payload["source_columns"])
        target_columns = self._safe_load_columns(payload["target_columns"])
        self._ensure_no_duplicate_relation(
            db,
            source_table=payload["source_table"],
            source_columns=source_columns,
            target_table=payload["target_table"],
            target_columns=target_columns,
            exclude_id=relation_id,
        )
        updated = Text2SQLTableRelationRepository(db).update(relation, payload)
        return self._to_relation_item(updated)

    def delete_relation(self, db: Session, relation_id: int) -> None:
        relation = self._get_owned_relation(db, relation_id)
        Text2SQLTableRelationRepository(db).delete(relation)

    def get_table_columns(self, db: Session, table_name: str) -> Text2SQLRelationTableColumnsResponse:
        real_table_name = self._resolve_table_name(db, table_name)
        table_detail = self.schema_service.get_table_detail(db, real_table_name)
        columns = [
            ColumnInfo(
                name=str(column.get("name") or ""),
                type=str(column.get("type") or ""),
                comment=str(column.get("comment") or ""),
            )
            for column in table_detail.get("columns", [])
            if str(column.get("name") or "").strip()
        ]
        return Text2SQLRelationTableColumnsResponse(
            table_name=real_table_name,
            table_comment=str(table_detail.get("table_comment") or ""),
            columns=columns,
        )

    def get_active_relations_by_tables(
        self,
        db: Session,
        table_names: list[str],
    ) -> list[dict]:
        resolved_tables, _ = self.schema_service.validate_selected_tables(db, table_names)
        if not resolved_tables:
            return []
        connection_key = self._normalize_connection_key(self.config_service.get_connection_key(db))
        rows = Text2SQLTableRelationRepository(db).list_by_tables(
            user_id=GLOBAL_CONFIG_USER_ID,
            connection_key=connection_key,
            table_names=resolved_tables,
        )
        relation_hints: list[dict] = []
        for item in rows:
            source_columns = self._safe_load_columns(item.source_columns)
            target_columns = self._safe_load_columns(item.target_columns)
            if not source_columns or len(source_columns) != len(target_columns):
                continue
            relation_hints.append(
                {
                    "id": int(item.id),
                    "source_table": str(item.source_table or ""),
                    "source_columns": source_columns,
                    "target_table": str(item.target_table or ""),
                    "target_columns": target_columns,
                    "relation_type": str(item.relation_type or ""),
                    "description": str(item.description or ""),
                    "summary": self._format_relation_summary(
                        str(item.source_table or ""),
                        source_columns,
                        str(item.target_table or ""),
                        target_columns,
                    ),
                }
            )
        return relation_hints

    @staticmethod
    def relation_hint_lines(relation_hints: list[dict] | None) -> list[str]:
        lines: list[str] = []
        for hint in relation_hints or []:
            source_table = str(hint.get("source_table") or "")
            target_table = str(hint.get("target_table") or "")
            source_columns = [str(item).strip() for item in (hint.get("source_columns") or []) if str(item).strip()]
            target_columns = [str(item).strip() for item in (hint.get("target_columns") or []) if str(item).strip()]
            if not source_table or not target_table or len(source_columns) != len(target_columns) or not source_columns:
                continue
            pairs = [
                f"{source_table}.{left} = {target_table}.{right}"
                for left, right in zip(source_columns, target_columns)
            ]
            relation_type = str(hint.get("relation_type") or "").strip()
            description = str(hint.get("description") or "").strip()
            suffix_parts = []
            if relation_type:
                suffix_parts.append(f"type: {relation_type}")
            if description:
                suffix_parts.append(f"description: {description}")
            suffix = f" ({'; '.join(suffix_parts)})" if suffix_parts else ""
            lines.append(f"- {' AND '.join(pairs)}{suffix}")
        return lines
