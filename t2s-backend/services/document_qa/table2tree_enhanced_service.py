from __future__ import annotations

import json
import re
from copy import deepcopy
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from json import JSONDecodeError
from typing import Any

import openpyxl
from langchain_core.prompts import ChatPromptTemplate
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from core.config import settings
from services.common.llm_service import common_llm_service

HEADER_ANALYSIS_PROMPT = """
你是一名专门解析复杂与非规则表格结构的资深数据分析专家。

输入表格使用了特殊的“单元格位置标注”格式，用以明确单元格的对齐与合并信息。
每个单元格以位置标识开头，例如 `A1`、`B2` 或 `A1:C1`。
- 示例：`A1:C1 招标说明` 表示内容“招标说明”跨越第 1 行的 A 到 C 列。
- 你必须利用这些位置信息，正确识别并对齐多行表头。
- 大多数合并单元格通常是表头。

你的任务是分析表格的表头行，并生成每列对应的标准化表头。

请严格遵守以下规则：

1. 严格保留原文（Strict wording preservation）：
   必须保留表头中的原始文字，不得改写、摘要、翻译或创造新名称。

2. 层级组合规则（Hierarchical combination）：
   如果表头跨越多行，则按 `[下层表头] - [上层表头]` 组合。

3. 精确列匹配（Exact column match）：
   输出数组长度必须与数据列数量完全一致，每个元素必须对应同一列。

4. 合并单元格规则：
   将 `A1:C1` 视为同时适用于 A、B、C 三列，不能只分配给左上角列。

[输入表格]:
{TABLE_AS_JSON_STRING}

你的输出必须是一个单一、有效的 JSON 字符串数组，代表最终的标准化表头。

[标准化表头]:
"""

HIERARCHY_VALUE_IDENTIFICATION_PROMPT = """
你是一名擅长分析表格逻辑结构的数据架构专家。

输入表格采用特殊 Markdown 格式，每个单元格都带有位置标识符，例如 "A1"、"B2:B4"。
请关注位置标识符后的实际内容，而不是标识符本身。

你将获得一个表格及其标准化表头。你的任务是识别哪些列属于“层次键”，哪些列属于“值叶节点”。

定义如下：
- 层次键（Hierarchy Keys）：用于分组和建立嵌套层级的列，其值通常在多个行块中重复出现。
- 值叶节点（Value Leaves）：与某个层次路径对应的最终数据列。
- 语义分组（Semantic Groups，可选）：如果值叶节点能根据上层表头归属于同一公共主题，请定义这些分组。
  组名必须直接来源于原始表格表头，不能创造新名称；不存在分组时必须返回空对象。

[输入表格]:
{TABLE_AS_JSON_STRING}

[标准化表头]:
{NORMALIZED_HEADERS_FROM_STEP_1}

你的输出必须是一个单一、有效的 JSON 对象，并严格使用原始表头名称：

{{
  "hierarchy_keys": ["header1", "header2"],
  "value_leaves": ["header3", "header4"],
  "semantic_groups": {{
    "原始分组名称1": ["header_a", "header_b"]
  }}
}}

不要提供任何解释性文字。

[层次定义]:
"""

TABLE_SUMMARY_PROMPT = """
你是一名数据目录与表格理解专家。请根据原始表格、标准化表头、层次键和值字段定义，为这张表生成一段面向用户的中文表格摘要。

输入表格使用 Markdown 样式，并且每个单元格带有位置标识符。位置标识符只用于理解表格结构，摘要中不要描述这些坐标。

请严格遵守：
1. 摘要要描述“这张表是什么、记录了哪些对象/主题、包含哪些主要维度和指标、适合回答什么类型的问题”。
2. 不要把标准化表头或层次定义原样复制成清单。
3. 不要输出 JSON、Markdown 表格、项目符号或解释过程。
4. 不要编造原始表格中没有的信息。
5. 控制在 120-220 个中文字符左右。

[输入表格]:
{TABLE_AS_JSON_STRING}

[标准化表头]:
{NORMALIZED_HEADERS_FROM_STEP_1}

[层次定义]:
{HIERARCHY_DEFINITION_FROM_STEP_2}

[表格摘要]:
"""

FINAL_JSON_TREE_CONSTRUCTION = """
你是一名高水平的数据转换引擎，能够将表格数据转换为嵌套的 JSON 树结构，并完整保留语义信息。

输入表格采用特殊 Markdown 格式，每个单元格都带有位置标识符：
- 单个单元格："A1"、"B2"、"C3"
- 合并单元格："A1:C1"、"A2:A4"

最终输出中请仅使用单元格的位置标识符作为值引用，而不是重复存储数据内容。

错误示例（禁止直接输出实际数值）：
{{
  "年级 - 1": {{
    "班级 - A": {{
      "学生总数": 30
    }}
  }}
}}

正确示例（叶值引用原始单元格坐标）：
{{
  "年级 - 1": {{
    "班级 - A": {{
      "学生总数": "C5"
    }}
  }}
}}

请严格遵守以下规则：

1. 语义层次键（Semantic Hierarchy Keys）：
   对每个 hierarchy_key，JSON 键必须格式化为 `[表头名称] - [单元格值]`，并使用空格-连字符-空格连接。

2. 叶节点键名严格保持：
   value_leaves 的键名必须与标准化表头完全一致，不能修改、缩写或翻译。

3. 值使用位置引用：
   所有叶值必须是单元格位置标识符，如 "A4"、"B12"、"C5:C7"，不能输出实际值。

4. 结构生成规则：
   遍历全部数据行并跳过表头；使用 hierarchy_keys 构建嵌套层级；在最深层写入 value_leaves。

5. 语义分组规则：
   若定义了 semantic_groups，使用来自原始表头的组名作为额外层级；不得自行生成分组名称。

6. 禁止数据遗漏：
   输入表格中的每个数据单元格位置都必须出现在最终树中，不得只输出示例、摘要、部分行或截断结构。

[输入表格]:
{TABLE_AS_JSON_STRING}

[标准化表头]:
{NORMALIZED_HEADERS_FROM_STEP_1}

[层次定义]:
{HIERARCHY_DEFINITION_FROM_STEP_2}

你的输出必须是一个单一、有效的 JSON 对象。不要输出任何解释性文字。

[最终 JSON 树]:
"""


@dataclass
class EnhancedTableParseResult:
    table_title: str
    markdown_table: str
    normalized_headers: str
    hierarchy_definition: str
    summary_text: str
    final_json_tree: str
    tree_with_cell_refs: dict[str, Any]
    tree: dict[str, Any]


class Table2TreeEnhancedService:
    def parse_sheet(self, sheet: Worksheet, *, db: Session | None = None) -> EnhancedTableParseResult:
        llm = common_llm_service.get_chat_model(db, temperature=0.2)
        if llm is None:
            raise ValueError("未配置 LLM，无法执行 enhanced table2tree 解析")

        table_title = extract_table_title(sheet)
        markdown_table = excel_to_markdown_with_cell_ref(sheet)
        normalized_headers = self._invoke_prompt(
            llm,
            HEADER_ANALYSIS_PROMPT,
            TABLE_AS_JSON_STRING=markdown_table,
        )
        hierarchy_definition = self._invoke_prompt(
            llm,
            HIERARCHY_VALUE_IDENTIFICATION_PROMPT,
            TABLE_AS_JSON_STRING=markdown_table,
            NORMALIZED_HEADERS_FROM_STEP_1=normalized_headers,
        )
        summary_text = self._invoke_prompt(
            llm,
            TABLE_SUMMARY_PROMPT,
            TABLE_AS_JSON_STRING=markdown_table,
            NORMALIZED_HEADERS_FROM_STEP_1=normalized_headers,
            HIERARCHY_DEFINITION_FROM_STEP_2=hierarchy_definition,
        )
        final_json_tree = self._invoke_prompt(
            llm,
            FINAL_JSON_TREE_CONSTRUCTION,
            TABLE_AS_JSON_STRING=markdown_table,
            NORMALIZED_HEADERS_FROM_STEP_1=normalized_headers,
            HIERARCHY_DEFINITION_FROM_STEP_2=hierarchy_definition,
        )
        tree_with_cell_refs = parse_json_with_merge(final_json_tree)
        tree = convert_cell_positions_to_values(tree_with_cell_refs, sheet)
        return EnhancedTableParseResult(
            table_title=table_title,
            markdown_table=markdown_table,
            normalized_headers=normalized_headers,
            hierarchy_definition=hierarchy_definition,
            summary_text=summary_text,
            final_json_tree=final_json_tree,
            tree_with_cell_refs=tree_with_cell_refs,
            tree=tree,
        )

    @staticmethod
    def _invoke_prompt(llm: Any, prompt_template: str, **kwargs: str) -> str:
        prompt = ChatPromptTemplate.from_template(prompt_template)
        response = (prompt | llm).invoke(kwargs)
        return _message_to_text(response)


def excel_to_markdown_with_cell_ref(sheet: Worksheet) -> str:
    merged_cells: dict[str, Any] = {}
    for merged_range in sheet.merged_cells.ranges:
        merged_ref = (
            f"{get_column_letter(merged_range.min_col)}{merged_range.min_row}:"
            f"{get_column_letter(merged_range.max_col)}{merged_range.max_row}"
        )
        merged_cells[merged_ref] = sheet.cell(merged_range.min_row, merged_range.min_col).value

    table_data: list[list[str]] = []
    for row in range(1, sheet.max_row + 1):
        row_data: list[str] = []
        for col in range(1, sheet.max_column + 1):
            cell = sheet.cell(row, col)
            cell_ref = f"{get_column_letter(col)}{row}"
            cell_value = None
            is_in_merged = False

            for merged_ref, value in merged_cells.items():
                start_ref, end_ref = merged_ref.split(":")
                start_col = openpyxl.utils.column_index_from_string("".join(filter(str.isalpha, start_ref)))
                start_row = int("".join(filter(str.isdigit, start_ref)))
                end_col = openpyxl.utils.column_index_from_string("".join(filter(str.isalpha, end_ref)))
                end_row = int("".join(filter(str.isdigit, end_ref)))
                if start_row <= row <= end_row and start_col <= col <= end_col:
                    is_in_merged = True
                    if row == start_row and col == start_col:
                        cell_value = f"{merged_ref} {value if value is not None else ''}"
                    break

            if not is_in_merged:
                cell_value = f"{cell_ref} {cell.value if cell.value is not None else ''}"
            row_data.append(cell_value if cell_value is not None else f"{cell_ref} ")
        table_data.append(row_data)

    if not table_data:
        return ""
    lines = [
        "| " + " | ".join(str(cell) for cell in table_data[0]) + " |",
        "| " + " | ".join(["---"] * sheet.max_column) + " |",
    ]
    lines.extend("| " + " | ".join(str(cell) for cell in row) + " |" for row in table_data[1:])
    return "\n".join(lines)


def extract_table_title(sheet: Worksheet, max_scan_rows: int = 5) -> str:
    for row in range(1, min(sheet.max_row, max_scan_rows) + 1):
        row_values: list[tuple[int, str]] = []
        for col in range(1, sheet.max_column + 1):
            text = _clean_cell_text(sheet.cell(row, col).value)
            if text:
                row_values.append((col, text))
        if not row_values:
            continue
        merged_title = _merged_row_title(sheet, row, row_values)
        if merged_title:
            return merged_title
        if row <= 2 and len(row_values) == 1:
            return row_values[0][1]
    return str(sheet.title)


def large_table_reason(sheet: Worksheet, markdown_table: str) -> str | None:
    reasons: list[str] = []
    cell_count = int(sheet.max_row or 0) * int(sheet.max_column or 0)
    merged_ranges = getattr(getattr(sheet, "merged_cells", None), "ranges", [])
    merged_count = len(merged_ranges)
    if cell_count > settings.LARGE_TABLE_CELL_THRESHOLD:
        reasons.append(f"cells={cell_count}>{settings.LARGE_TABLE_CELL_THRESHOLD}")
    if int(sheet.max_row or 0) > settings.LARGE_TABLE_ROW_THRESHOLD:
        reasons.append(f"rows={sheet.max_row}>{settings.LARGE_TABLE_ROW_THRESHOLD}")
    if int(sheet.max_column or 0) > settings.LARGE_TABLE_COLUMN_THRESHOLD:
        reasons.append(f"columns={sheet.max_column}>{settings.LARGE_TABLE_COLUMN_THRESHOLD}")
    if merged_count > settings.LARGE_TABLE_MERGED_CELL_THRESHOLD:
        reasons.append(f"merged={merged_count}>{settings.LARGE_TABLE_MERGED_CELL_THRESHOLD}")
    if len(markdown_table) > settings.LARGE_TABLE_MARKDOWN_THRESHOLD:
        reasons.append(f"markdown={len(markdown_table)}>{settings.LARGE_TABLE_MARKDOWN_THRESHOLD}")
    return "; ".join(reasons) if reasons else None


def validate_enhanced_tree_quality(tree_with_cell_refs: dict[str, Any]) -> list[str]:
    warnings: list[str] = []
    leaf_count = 0
    bad_leaf_count = 0
    bad_examples: list[str] = []
    suspicious_examples: list[str] = []
    suspicious_patterns = ["极简输出", "仅展示", "完整输出", "如需", "结构示例", "模式已启用", "省略", "未完", "truncated"]

    def walk(value: Any, path: list[str]) -> None:
        nonlocal leaf_count, bad_leaf_count
        if isinstance(value, dict):
            for key, item in value.items():
                walk(item, [*path, str(key)])
            return
        if isinstance(value, list):
            for index, item in enumerate(value):
                walk(item, [*path, str(index)])
            return

        leaf_count += 1
        text = str(value)
        if any(pattern in text for pattern in suspicious_patterns) and len(suspicious_examples) < 3:
            suspicious_examples.append(f"{' | '.join(path)}={text[:120]}")
        if not _is_cell_position(text) and not _is_cell_range(text):
            bad_leaf_count += 1
            if len(bad_examples) < 3:
                bad_examples.append(f"{' | '.join(path)}={text[:120]}")

    walk(tree_with_cell_refs, [])
    if leaf_count == 0:
        warnings.append("no_leaf_values")
    if suspicious_examples:
        warnings.append("model_comment_in_tree: " + " || ".join(suspicious_examples))
    if bad_leaf_count:
        warnings.append(f"non_cell_ref_leaf_values={bad_leaf_count}/{leaf_count}: " + " || ".join(bad_examples))
    return warnings


def parse_json_with_merge(llm_output: str) -> dict[str, Any]:
    json_content = _extract_json_object(llm_output)

    def deep_merge(left: dict[str, Any], right: dict[str, Any]) -> dict[str, Any]:
        merged = left.copy()
        for key, right_value in right.items():
            if key not in merged:
                merged[key] = right_value
                continue
            left_value = merged[key]
            if isinstance(left_value, dict) and isinstance(right_value, dict):
                merged[key] = deep_merge(left_value, right_value)
            else:
                if not isinstance(left_value, list):
                    merged[key] = [left_value]
                if isinstance(right_value, list):
                    merged[key].extend(right_value)
                else:
                    merged[key].append(right_value)
        return merged

    def recursive_merge_hook(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
        result: dict[str, Any] = {}
        for key, value in pairs:
            if key in result:
                if isinstance(result[key], dict) and isinstance(value, dict):
                    result[key] = deep_merge(result[key], value)
                else:
                    if not isinstance(result[key], list):
                        result[key] = [result[key]]
                    if isinstance(value, list):
                        result[key].extend(value)
                    else:
                        result[key].append(value)
            else:
                result[key] = value
        return result

    try:
        parsed = json.loads(json_content, object_pairs_hook=recursive_merge_hook)
    except JSONDecodeError as first_exc:
        repaired_json = _repair_common_json_issues(json_content)
        try:
            parsed = json.loads(repaired_json, object_pairs_hook=recursive_merge_hook)
        except JSONDecodeError:
            raise first_exc
    if not isinstance(parsed, dict):
        raise ValueError("LLM final tree output must be a JSON object")
    return parsed


def convert_cell_positions_to_values(data_dict: dict[str, Any], sheet: Worksheet, range_policy: str = "merged_top_left") -> dict[str, Any]:
    result = deepcopy(data_dict)

    def recursive_convert(obj: Any) -> Any:
        if isinstance(obj, dict):
            return {key: recursive_convert(value) for key, value in obj.items()}
        if isinstance(obj, list):
            return [recursive_convert(item) for item in obj]
        if isinstance(obj, str):
            return get_value_from_ref(obj, sheet, range_policy=range_policy)
        return obj

    return recursive_convert(result)


def get_value_from_ref(ref: str, sheet: Worksheet, range_policy: str) -> Any:
    ref_u = ref.strip().upper()
    if _is_cell_position(ref_u):
        try:
            return _json_safe(sheet[ref_u].value)
        except Exception:
            return ref
    if _is_cell_range(ref_u):
        try:
            min_col, min_row, max_col, max_row = range_boundaries(ref_u)
            for merged_range in sheet.merged_cells.ranges:
                if range_boundaries(str(merged_range).upper()) == (min_col, min_row, max_col, max_row):
                    return _json_safe(sheet.cell(row=min_row, column=min_col).value)
            cells = sheet[ref_u]
            if range_policy == "matrix":
                if isinstance(cells, tuple) and cells and isinstance(cells[0], tuple):
                    return [[_json_safe(cell.value) for cell in row] for row in cells]
                return [_json_safe(cell.value) for cell in cells]
            if range_policy == "flatten":
                if isinstance(cells, tuple) and cells and isinstance(cells[0], tuple):
                    return [_json_safe(cell.value) for row in cells for cell in row]
                return [_json_safe(cell.value) for cell in cells]
            return _json_safe(sheet.cell(row=min_row, column=min_col).value)
        except Exception:
            return ref
    return ref


def _merged_row_title(sheet: Worksheet, row: int, row_values: list[tuple[int, str]]) -> str:
    min_title_span = max(2, (sheet.max_column * 3 + 4) // 5)
    for col, text in row_values:
        for merged_range in sheet.merged_cells.ranges:
            if (
                merged_range.min_row <= row <= merged_range.max_row
                and merged_range.min_col <= col <= merged_range.max_col
                and merged_range.max_col - merged_range.min_col + 1 >= min_title_span
            ):
                return text
    return ""


def _clean_cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _message_to_text(message: Any) -> str:
    content = getattr(message, "content", message)
    if isinstance(content, str):
        return _extract_after_think(content)
    if isinstance(content, list):
        parts = []
        for item in content:
            if isinstance(item, dict) and item.get("type") == "text":
                parts.append(str(item.get("text", "")))
            else:
                parts.append(str(item))
        return _extract_after_think("".join(parts))
    return _extract_after_think(str(content))


def _extract_after_think(response: str) -> str:
    if "</think>" in response:
        return response.split("</think>", 1)[1].strip()
    return response.strip()


def _extract_json_object(llm_output: str) -> str:
    if "```json" in llm_output:
        return llm_output.split("```json", 1)[1].split("```", 1)[0].strip()
    if "```" in llm_output:
        return llm_output.split("```", 1)[1].split("```", 1)[0].strip()
    start = llm_output.find("{")
    end = llm_output.rfind("}") + 1
    if start < 0 or end <= start:
        raise ValueError("No JSON object found in LLM output")
    return llm_output[start:end].strip()


def _repair_common_json_issues(json_content: str) -> str:
    repaired = re.sub(r",(\s*[}\]])", r"\1", json_content)
    start = repaired.find("{")
    end = repaired.rfind("}") + 1
    if start >= 0 and end > start:
        repaired = repaired[start:end]
    return repaired.strip()


def _is_cell_position(text: str) -> bool:
    return bool(re.match(r"^\$?[A-Z]+\$?[0-9]+$", text.strip().upper()))


def _is_cell_range(text: str) -> bool:
    return bool(re.match(r"^\$?[A-Z]+\$?[0-9]+:\$?[A-Z]+\$?[0-9]+$", text.strip().upper()))


def _json_safe(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, list):
        return [_json_safe(item) for item in value]
    if isinstance(value, tuple):
        return [_json_safe(item) for item in value]
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


table2tree_enhanced_service = Table2TreeEnhancedService()
