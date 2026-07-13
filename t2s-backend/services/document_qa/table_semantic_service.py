from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import math
import os
import re
from collections.abc import Callable
from dataclasses import asdict, dataclass
from datetime import date, datetime, time
from typing import Any
from uuid import uuid4

from fastapi import UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from core.config import settings
from core.knowledge_usage import KB_USAGE_TABLE_SEMANTIC_TREE
from core.domain_errors import (
    FileProcessorMismatchError,
    FileUsageSnapshotMismatchError,
    TableArtifactContextMismatchError,
    TableArtifactNotFoundError,
)
from core.knowledge_policy import (
    KnowledgeOperation,
    ProcessorType,
)
from services.common.knowledge_guard_service import (
    knowledge_guard,
)
from models.document_chunk import DocumentChunk
from models.knowledge_file import KnowledgeFile
from models.table_semantic_artifact import TableSemanticArtifact
from repositories.es_repo import es_repo
from repositories.knowledge_base_repo import KnowledgeBaseRepository
from repositories.knowledge_file_repo import KnowledgeFileRepository
from repositories.minio_repo import minio_repo
from repositories.table_semantic_artifact_repo import TableSemanticArtifactRepository
from services.common.embeddings import get_embedding_vector_dim, get_embeddings
from services.document_qa.table_parse_plan_service import PlanBuildResult, table_parse_plan_service
from services.document_qa.table_search_index_service import build_tree_index_fields, table_search_index_service
from services.document_qa.table2tree_enhanced_service import (
    EnhancedTableParseResult,
    excel_to_markdown_with_cell_ref,
    large_table_reason,
    table2tree_enhanced_service,
    validate_enhanced_tree_quality,
)

MAX_TREE_ROWS = 1000
MAX_TREE_PATHS = 4000
PATHS_PER_CHUNK = 40
SUPPORTED_TABLE_TYPES = {"xlsx", "xlsm", "xls", "csv"}
logger = logging.getLogger(__name__)


@dataclass
class ParsedTable:
    sheet_name: str
    title: str
    summary_text: str
    headers: list[str]
    records: list[dict[str, Any]]
    tree: dict[str, Any]
    tree_with_cell_refs: dict[str, Any]
    tree_path_text: list[str]
    tree_metric_names: list[str]
    row_count: int
    column_count: int
    parse_plan: dict[str, Any] | None = None
    parse_mode: str = "heuristic"
    large_table_reason: str = ""
    markdown_table: str = ""
    normalized_headers: str = ""
    hierarchy_definition: str = ""
    final_json_tree: str = ""
    coverage: dict[str, Any] | None = None
    validation_warnings: list[str] | None = None


def _sanitize_filename(value: str) -> str:
    text = os.path.basename(str(value or "").strip())
    text = re.sub(r"[\\/:*?\"<>|]", "_", text)
    return text or "unnamed_table"


def _safe_json(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(key): _safe_json(item) for key, item in value.items()}
    if isinstance(value, list):
        return [_safe_json(item) for item in value]
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def _is_empty(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value).strip()


class TableSemanticService:
    async def upload_table(
        self,
        *,
        db: Session,
        file: UploadFile,
        kb_id: int | None = None,
        sheet_name: str | None = None,
    ) -> dict:
        filename = str(file.filename or "").strip()
        if not filename:
            raise ValueError("Uploaded file must have a filename")

        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        if ext not in SUPPORTED_TABLE_TYPES:
            raise ValueError(f"Unsupported table file type: .{ext}")

        kb = self._resolve_table_kb(db, kb_id)
        raw = await file.read()
        await file.seek(0)

        max_bytes = int(settings.MAX_UPLOAD_FILE_SIZE_MB) * 1024 * 1024
        if len(raw) > max_bytes:
            raise ValueError(f"File exceeds size limit: {settings.MAX_UPLOAD_FILE_SIZE_MB}MB")

        md5 = hashlib.md5(raw, usedforsecurity=False).hexdigest()
        batch_id = uuid4().hex
        celery_task_id = uuid4().hex
        sheet_jobs = self._build_sheet_jobs(raw=raw, file_type=ext, sheet_name=sheet_name)

        table_ids = [
            str(item["table_id"])
            for item in sheet_jobs
        ]

        sheet_names = [
            str(item.get("sheet_name") or "")
            for item in sheet_jobs
        ]

        initial_task_meta = {
            "file_id": None,
            "batch_id": batch_id,
            "state": "PENDING",

            "table_ids": table_ids,
            "sheet_names": sheet_names,

            "total_sheets": len(sheet_jobs),
            "completed_sheets": 0,
            "progress": 0.0,

            "completed_table_ids": [],
            "current_table_id": None,
            "current_sheet_name": None,

            "error": None,
        }

        safe_name = _sanitize_filename(filename)
        object_name = f"table_semantic/kb_{int(kb.id)}/{batch_id}/source/{safe_name}"
        minio_repo.upload_file_bytes(
            object_name,
            raw,
            content_type=file.content_type or "application/octet-stream",
        )

        file_repo = KnowledgeFileRepository(db)
        from core.knowledge_policy import ProcessorType
        from core.knowledge_usage import KB_USAGE_TABLE_SEMANTIC_TREE
        entity = KnowledgeFile(
            kb_id=int(kb.id),
            file_name=safe_name,
            file_type=ext,
            file_size=len(raw),
            md5=md5,
            minio_bucket=minio_repo.bucket_name,
            minio_object_name=object_name,
            status=0,
            task_id=celery_task_id,

            table_task_meta_json=json.dumps(
                initial_task_meta,
                ensure_ascii=False,
            ),

            is_deleted=False,
            usage_snapshot=(
                KB_USAGE_TABLE_SEMANTIC_TREE
            ),
            processor_type=(
                ProcessorType.TABLE_SEMANTIC.value
            ),
            process_version=1,
        )

        created = file_repo.create_file(entity)

        initial_task_meta["file_id"] = int(
            created.id
        )

        file_repo.update_table_task_meta_for_task(
            int(created.id),
            task_id=celery_task_id,
            process_version=1,
            task_meta=initial_task_meta,
        )

        from tasks.table_semantic_tasks import process_table_semantic_task

        try:
            process_table_semantic_task.apply_async(
                kwargs={
                    "file_id": int(created.id),
                    "batch_id": batch_id,
                    "sheet_jobs": sheet_jobs,
                    "expected_usage": (
                        KB_USAGE_TABLE_SEMANTIC_TREE
                    ),
                    "expected_processor_type": (
                        ProcessorType.TABLE_SEMANTIC.value
                    ),
                    "process_version": 1,
                },
                task_id=celery_task_id,
            )
        except Exception as exc:
            file_repo.update_status_for_task(
                int(created.id),
                task_id=celery_task_id,
                process_version=1,
                status=3,
                error_msg=(
                    "Failed to publish table task: "
                    f"{exc}"
                )[:1000],
            )
            raise
        return {
            "file_id": int(created.id),
            "task_id": celery_task_id,
            "batch_id": batch_id,
            "table_id": table_ids[0],
            "table_ids": table_ids,
            "sheet_names": sheet_names,
            "status": "queued",
            "message": (
                f"表格解析任务已提交，"
                f"共 {len(sheet_jobs)} 张工作表"
            ),
        }

    def process_table_file(
        self,
        db: Session,
        *,
        file_entity: KnowledgeFile,
        table_id: str,
        sheet_name: str | None = None,
    ) -> dict:
        return self.process_table_batch(
            db,
            file_entity=file_entity,
            batch_id=str(table_id),
            sheet_jobs=[{"table_id": str(table_id), "sheet_name": sheet_name}],
        )

    def process_table_batch(
        self,
        db: Session,
        *,
        file_entity: KnowledgeFile,
        batch_id: str,
        sheet_jobs: list[dict[str, Any]],
        progress_callback: Callable[
            [dict[str, Any]],
            None,
        ] | None = None,
    ) -> dict:
        guard_ctx = (
            knowledge_guard.resolve_file_for_operation(
                db,
                int(file_entity.id),
                KnowledgeOperation.TABLE_REPROCESS,
            )
        )

        guard_file = guard_ctx.file
        kb = guard_ctx.kb

        if guard_file is None:
            raise ValueError("Source file not found")

        if (
            int(guard_file.id)
            != int(file_entity.id)
        ):
            raise TableArtifactContextMismatchError(
                "Worker file context changed"
            )

        if (
            guard_ctx.processor_type
            != ProcessorType.TABLE_SEMANTIC
        ):
            raise FileProcessorMismatchError(
                "Table processing requires "
                "table_semantic processor"
            )

        raw = minio_repo.get_file_bytes(file_entity.minio_object_name)
        file_repo = KnowledgeFileRepository(db)
        artifact_repo = TableSemanticArtifactRepository(db)
        
        file_repo.delete_chunks_by_file_id(
            int(file_entity.id),
            kb_id=int(kb.id),
            usage_snapshot=(
                KB_USAGE_TABLE_SEMANTIC_TREE
            ),
        )

        artifact_repo.soft_delete_by_file_context(
            file_id=int(file_entity.id),
            kb_id=int(kb.id),
        )
        try:
            es_repo.delete_chunks_by_file_id(int(file_entity.id), index_name=str(kb.collection_name))
        except Exception:
            pass
        try:
            table_search_index_service.delete_by_file_context(
                collection_name=str(
                    kb.collection_name
                ),
                kb_id=int(kb.id),
                file_id=int(file_entity.id),
            )
        except Exception as exc:  # noqa: BLE001
            logger.warning("delete table search index by file failed: file_id=%s error=%s", file_entity.id, exc)

        tables: list[dict[str, Any]] = []
        total_chunk_count = 0

        total_sheets = len(sheet_jobs)

        planned_table_ids = [
            str(item.get("table_id") or "")
            for item in sheet_jobs
        ]

        planned_sheet_names = [
            str(item.get("sheet_name") or "")
            for item in sheet_jobs
        ]

        completed_table_ids: list[str] = []

        for index, job in enumerate(sheet_jobs, start=1):
            table_id = str(
                job.get("table_id") or uuid4().hex
            )
            sheet_name = (
                str(job.get("sheet_name") or "").strip()
                or None
            )

            if progress_callback is not None:
                progress_callback(
                    {
                        "total_sheets": total_sheets,
                        "completed_sheets": index - 1,
                        "progress": (
                            (index - 1) / total_sheets
                            if total_sheets
                            else 0.0
                        ),
                        "table_ids": planned_table_ids,
                        "sheet_names": planned_sheet_names,
                        "completed_table_ids": list(
                            completed_table_ids
                        ),
                        "current_table_id": table_id,
                        "current_sheet_name": sheet_name,
                    }
                )

            parsed = self.parse_table(
                raw=raw,
                file_type=str(file_entity.file_type),
                file_name=str(file_entity.file_name),
                sheet_name=sheet_name,
                db=db,
                use_enhanced=True,
            )

            normalized_object_name = (
                f"table_semantic/kb_{int(kb.id)}/"
                f"{batch_id}/{table_id}/normalized.xlsx"
            )

            minio_repo.upload_file_bytes(
                normalized_object_name,
                self._build_normalized_workbook(parsed),
                content_type=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
            )

            result = self._store_parsed_table(
                db,
                file_entity=file_entity,
                collection_name=str(kb.collection_name),
                batch_id=str(batch_id),
                table_id=table_id,
                parsed=parsed,
                normalized_object_name=normalized_object_name,
            )

            total_chunk_count += int(result["chunk_count"])
            tables.append(result)
            completed_table_ids.append(table_id)

            if progress_callback is not None:
                progress_callback(
                    {
                        "total_sheets": total_sheets,
                        "completed_sheets": index,
                        "progress": (
                            index / total_sheets
                            if total_sheets
                            else 1.0
                        ),
                        "table_ids": planned_table_ids,
                        "sheet_names": planned_sheet_names,
                        "completed_table_ids": list(
                            completed_table_ids
                        ),
                        "current_table_id": None,
                        "current_sheet_name": None,
                    }
                )

        return {
            "file_id": int(file_entity.id),
            "kb_id": int(file_entity.kb_id),
            "batch_id": str(batch_id),

            "table_ids": planned_table_ids,
            "sheet_names": planned_sheet_names,

            "total_sheets": total_sheets,
            "completed_sheets": total_sheets,
            "progress": 1.0,

            "current_table_id": None,
            "current_sheet_name": None,

            "tables": tables,
            "chunk_count": total_chunk_count,
        }

    def _store_parsed_table(
        self,
        db: Session,
        *,
        file_entity: KnowledgeFile,
        collection_name: str,
        batch_id: str,
        table_id: str,
        parsed: ParsedTable,
        normalized_object_name: str,
    ) -> dict:
        file_repo = KnowledgeFileRepository(db)
        artifact_repo = TableSemanticArtifactRepository(db)
        tree_object_name = f"table_semantic/kb_{int(file_entity.kb_id)}/{batch_id}/{table_id}/tree.json"
        tree_payload = {
            "batch_id": batch_id,
            "table_id": table_id,
            "source_object_name": str(file_entity.minio_object_name),
            "normalized_object_name": normalized_object_name,
            "summary_text": parsed.summary_text,
            "candidate_fields": _safe_json(parsed.headers),
            "tree_path_text": _safe_json(parsed.tree_path_text[:MAX_TREE_PATHS]),
            "tree_metric_names": _safe_json(parsed.tree_metric_names),
            "tree": _safe_json(parsed.tree),
            "tree_with_cell_refs": _safe_json(parsed.tree_with_cell_refs),
            "parse_plan": _safe_json(parsed.parse_plan or {}),
            "parse_mode": parsed.parse_mode or _parse_mode(parsed.parse_plan),
            "large_table_reason": parsed.large_table_reason,
            "markdown_table": parsed.markdown_table,
            "normalized_headers": parsed.normalized_headers,
            "hierarchy_definition": parsed.hierarchy_definition,
            "final_json_tree": parsed.final_json_tree,
            "coverage": _safe_json(parsed.coverage or {}),
            "validation_warnings": _safe_json(parsed.validation_warnings or []),
        }
        minio_repo.upload_file_bytes(
            tree_object_name,
            json.dumps(
                tree_payload,
                ensure_ascii=False,
            ).encode("utf-8"),
            content_type="application/json",
        )

        chunk_entities = [
            DocumentChunk(
                kb_id=int(file_entity.kb_id),
                file_id=int(file_entity.id),
                chunk_index=index,
                content=content,
                char_count=len(content),
                usage_snapshot=file_entity.usage_snapshot,
                processor_type=file_entity.processor_type,
            )
            for index, content in enumerate(self._build_index_chunks(table_id, parsed))
        ]
        saved_chunks = file_repo.bulk_create_chunks(chunk_entities)
        self._index_chunks(db, collection_name, saved_chunks)

        artifact = TableSemanticArtifact(
            kb_id=int(file_entity.kb_id),
            file_id=int(file_entity.id),
            table_id=table_id,
            file_name=str(file_entity.file_name),
            sheet_name=parsed.sheet_name,
            table_title=parsed.title,
            summary_text=parsed.summary_text,
            candidate_fields_json=json.dumps(parsed.headers, ensure_ascii=False),
            tree_path_text_json=json.dumps(parsed.tree_path_text[:MAX_TREE_PATHS], ensure_ascii=False),
            tree_metric_names_json=json.dumps(parsed.tree_metric_names, ensure_ascii=False),
            tree_object_name=tree_object_name,
            row_count=parsed.row_count,
            column_count=parsed.column_count,
            usage_snapshot=file_entity.usage_snapshot,
        )
        created_artifact = artifact_repo.upsert(artifact)
        try:
            table_search_index_service.index_table(
                db,
                collection_name=collection_name,
                artifact=self.artifact_to_response(created_artifact),
                tree_payload=tree_payload,
            )
        except Exception as exc:  # noqa: BLE001
            logger.warning("index table search document failed: table_id=%s error=%s", table_id, exc)

        return {
            "file_id": int(file_entity.id),
            "kb_id": int(file_entity.kb_id),
            "batch_id": batch_id,
            "table_id": table_id,
            "sheet_name": parsed.sheet_name,
            "table_title": parsed.title,
            "chunk_count": len(saved_chunks),
            "row_count": parsed.row_count,
            "column_count": parsed.column_count,
            "normalized_object_name": normalized_object_name,
        }

    def parse_table(
        self,
        *,
        raw: bytes,
        file_type: str,
        file_name: str,
        sheet_name: str | None = None,
        db: Session | None = None,
        use_llm_plan: bool = False,
        use_enhanced: bool = False,
    ) -> ParsedTable:
        ext = str(file_type or "").lower().strip(".")
        if ext in {"xlsx", "xlsm"}:
            workbook = load_workbook(io.BytesIO(raw), data_only=True, read_only=False)
            if sheet_name and sheet_name not in workbook.sheetnames:
                raise ValueError(f"工作表不存在：{sheet_name}")
            selected_sheet = sheet_name or workbook.sheetnames[0]
            sheet = workbook[selected_sheet]
            if use_enhanced:
                markdown_table = excel_to_markdown_with_cell_ref(sheet)
                fallback_reason = large_table_reason(sheet, markdown_table)
                if not fallback_reason:
                    try:
                        enhanced_result = table2tree_enhanced_service.parse_sheet(sheet, db=db)
                        quality_warnings = validate_enhanced_tree_quality(enhanced_result.tree_with_cell_refs)
                        if not quality_warnings:
                            return self._from_enhanced_result(
                                result=enhanced_result,
                                sheet=sheet,
                                file_name=file_name,
                                markdown_table=markdown_table,
                            )
                        fallback_reason = "enhanced_tree_quality_failed: " + "; ".join(quality_warnings[:5])
                    except Exception as exc:  # noqa: BLE001
                        fallback_reason = f"enhanced_llm_failed: {exc}"

                result = table_parse_plan_service.parse_sheet_with_llm(
                    sheet,
                    file_name=file_name,
                    db=db,
                )
                parsed = self._from_plan_result(result=result, file_name=file_name)
                parsed.parse_mode = "plan_based"
                parsed.large_table_reason = fallback_reason or ""
                parsed.markdown_table = markdown_table
                parsed.normalized_headers = json.dumps(
                    {
                        "hierarchy_columns": parsed.parse_plan.get("hierarchy_columns", []) if parsed.parse_plan else [],
                        "value_columns": parsed.parse_plan.get("value_columns", parsed.headers) if parsed.parse_plan else parsed.headers,
                    },
                    ensure_ascii=False,
                    default=str,
                )
                parsed.hierarchy_definition = json.dumps(parsed.parse_plan or {}, ensure_ascii=False, default=str)
                warnings = list(parsed.validation_warnings or [])
                if fallback_reason:
                    warnings.append(fallback_reason)
                parsed.validation_warnings = warnings
                return parsed

            if use_llm_plan:
                result = table_parse_plan_service.parse_sheet_with_llm(
                    sheet,
                    file_name=file_name,
                    db=db,
                )
            else:
                result = table_parse_plan_service.parse_sheet(sheet, file_name=file_name)
            if not result.records:
                raise ValueError("未解析到有效数据行")
            return self._from_plan_result(result=result, file_name=file_name)

        rows, resolved_sheet = self._load_rows(raw, file_type=file_type, sheet_name=sheet_name)
        rows = self._trim_grid(rows)
        if not rows:
            raise ValueError("表格内容为空")

        header_index, title = self._detect_header(rows, fallback_title=file_name)
        headers = self._normalize_headers(rows[header_index])
        data_rows = rows[header_index + 1 :]
        records = self._build_records(headers, data_rows)
        if not records:
            raise ValueError("未解析到有效数据行")

        row_count = len(records)
        column_count = len(headers)
        tree_records = records[:MAX_TREE_ROWS]
        summary_text = self._build_summary(
            title=title,
            file_name=file_name,
            sheet_name=resolved_sheet,
            row_count=row_count,
            headers=headers,
        )
        tree_path_text = self._build_tree_paths(title=title, records=tree_records, limit=MAX_TREE_PATHS)
        tree_metric_names = headers[:120]
        tree = {
            "表格名": title,
            "摘要": summary_text,
            "字段": headers,
            "记录": tree_records,
        }
        tree_with_cell_refs = {
            "表格名": title,
            "摘要": summary_text,
            "字段": headers,
            "路径": tree_path_text,
        }
        return ParsedTable(
            sheet_name=resolved_sheet,
            title=title,
            summary_text=summary_text,
            headers=headers,
            records=records,
            tree=tree,
            tree_with_cell_refs=tree_with_cell_refs,
            tree_path_text=tree_path_text,
            tree_metric_names=tree_metric_names,
            row_count=row_count,
            column_count=column_count,
            parse_plan=None,
            coverage=None,
            validation_warnings=[],
        )

    def _resolve_table_context(
        self,
        db: Session,
        table_id: str,
    ):
        normalized_table_id = str(
            table_id or ""
        ).strip()

        if not normalized_table_id:
            raise TableArtifactNotFoundError(
                "table_id cannot be empty"
            )

        artifact_repo = (
            TableSemanticArtifactRepository(db)
        )

        raw_artifact = (
            artifact_repo.get_any_by_table_id(
                normalized_table_id
            )
        )

        if raw_artifact is None:
            raise TableArtifactNotFoundError(
                f"Table artifact not found: "
                f"{normalized_table_id}"
            )

        row = artifact_repo.get_with_context(
            normalized_table_id
        )

        if row is None:
            raise TableArtifactContextMismatchError(
                "Table artifact exists but its KB, file, "
                "usage snapshot or processor context "
                "is invalid"
            )

        artifact, file_entity, kb_entity = row

        guard_ctx = (
            knowledge_guard.resolve_file_for_operation(
                db,
                int(file_entity.id),
                KnowledgeOperation.TABLE_READ,
            )
        )

        if (
            guard_ctx.processor_type
            != ProcessorType.TABLE_SEMANTIC
        ):
            raise FileProcessorMismatchError(
                "Table artifact source file processor "
                "must be table_semantic"
            )

        if (
            int(artifact.kb_id)
            != int(guard_ctx.kb.id)
            or int(artifact.file_id)
            != int(file_entity.id)
        ):
            raise TableArtifactContextMismatchError(
                "Table artifact identifiers do not match "
                "the guarded file context"
            )

        if (
            str(artifact.usage_snapshot)
            != guard_ctx.usage.value
        ):
            raise FileUsageSnapshotMismatchError(
                "Table artifact usage snapshot does not "
                "match its knowledge base usage"
            )

        return (
            artifact,
            file_entity,
            kb_entity,
            guard_ctx,
        )

    def get_table(
        self,
        db: Session,
        table_id: str,
    ) -> TableSemanticArtifact:
        artifact, _, _, _ = (
            self._resolve_table_context(
                db,
                table_id,
            )
        )

        return artifact

    def list_tables(
        self,
        db: Session,
        *,
        kb_id: int | None = None,
        limit: int = 50,
    ) -> list[TableSemanticArtifact]:
        artifact_repo = (
            TableSemanticArtifactRepository(db)
        )

        if kb_id is not None:
            guard_ctx = (
                knowledge_guard.resolve_kb_for_operation(
                    db,
                    int(kb_id),
                    KnowledgeOperation.TABLE_READ,
                )
            )

            return artifact_repo.list_valid_by_kb_ids(
                kb_ids=[int(guard_ctx.kb.id)],
                limit=limit,
            )

        kb_rows = (
            KnowledgeBaseRepository(db)
            .list_by_usage(
                KB_USAGE_TABLE_SEMANTIC_TREE
            )
        )

        kb_ids = [
            int(item.id)
            for item in kb_rows
        ]

        return artifact_repo.list_valid_by_kb_ids(
            kb_ids=kb_ids,
            limit=limit,
        )

    def get_table_tree(self, db: Session, table_id: str) -> dict:
        entity = self.get_table(db, table_id)
        raw = minio_repo.get_file_bytes(str(entity.tree_object_name))
        return json.loads(raw.decode("utf-8"))

    def get_table_artifact(self, db: Session, table_id: str) -> dict:
        entity = self.get_table(db, table_id)
        payload = self.artifact_to_response(entity)
        try:
            tree_payload = self.get_table_tree(db, table_id)
        except Exception:
            tree_payload = {}
        payload.update(
            {
                "batch_id": tree_payload.get("batch_id") or "",
                "source_object_name": tree_payload.get("source_object_name") or "",
                "normalized_object_name": tree_payload.get("normalized_object_name") or "",
                "tree": tree_payload.get("tree") or {},
                "tree_with_cell_refs": tree_payload.get("tree_with_cell_refs") or {},
                "parse_plan": tree_payload.get("parse_plan") or {},
                "parse_mode": tree_payload.get("parse_mode") or _parse_mode(tree_payload.get("parse_plan")),
                "large_table_reason": tree_payload.get("large_table_reason") or "",
                "markdown_table": tree_payload.get("markdown_table") or "",
                "normalized_headers": tree_payload.get("normalized_headers") or "",
                "hierarchy_definition": tree_payload.get("hierarchy_definition") or "",
                "final_json_tree": tree_payload.get("final_json_tree") or "",
                "coverage": tree_payload.get("coverage") or {},
                "validation_warnings": tree_payload.get("validation_warnings") or [],
            }
        )
        return payload

    def get_table_source_url(self, db: Session, table_id: str) -> dict:
        entity = self.get_table(db, table_id)
        file_entity = KnowledgeFileRepository(db).get_by_id(int(entity.file_id))
        if file_entity is None:
            raise ValueError("Source file not found")
        return {
            "table_id": str(entity.table_id),
            "file_id": int(entity.file_id),
            "file_name": str(entity.file_name),
            "url": minio_repo.get_presigned_url(str(file_entity.minio_object_name)),
        }

    def get_table_normalized_url(self, db: Session, table_id: str) -> dict:
        entity = self.get_table(db, table_id)
        tree_payload = self.get_table_tree(db, table_id)
        object_name = str(tree_payload.get("normalized_object_name") or "").strip()
        if not object_name:
            raise ValueError("Normalized file not found")
        file_name = f"{_sanitize_filename(entity.table_title)}.xlsx"
        return {
            "table_id": str(entity.table_id),
            "file_id": int(entity.file_id),
            "file_name": file_name,
            "url": minio_repo.get_presigned_url(object_name),
        }

    def get_table_index_document(self, db: Session, table_id: str) -> dict[str, Any]:
        entity = self.get_table(db, table_id)
        kb = KnowledgeBaseRepository(db).get_by_id(int(entity.kb_id))
        if kb is None:
            raise ValueError("Knowledge base not found")
        document = table_search_index_service.get_table_document(
            collection_name=str(kb.collection_name),
            table_id=str(table_id),
        )
        if document:
            valid_document = (
                int(document.get("kb_id") or 0)
                == int(entity.kb_id)
                and int(document.get("file_id") or 0)
                == int(entity.file_id)
                and str(
                    document.get("usage_type") or ""
                )
                == KB_USAGE_TABLE_SEMANTIC_TREE
                and str(
                    document.get("processor_type") or ""
                )
                == ProcessorType.TABLE_SEMANTIC.value
                and str(
                    document.get("table_id") or ""
                )
                == str(entity.table_id)
            )

            if valid_document:
                return document

            logger.error(
                "Invalid table index document context: "
                "table_id=%s kb_id=%s file_id=%s",
                entity.table_id,
                entity.kb_id,
                entity.file_id,
            )
        artifact = self.get_table_artifact(db, table_id)
        tree_fields = build_tree_index_fields(artifact.get("tree") if isinstance(artifact.get("tree"), dict) else {})
        parse_plan = artifact.get("parse_plan") if isinstance(artifact.get("parse_plan"), dict) else {}
        normalized_headers = str(
            artifact.get("normalized_headers")
            or json.dumps(
                {
                    "hierarchy_columns": parse_plan.get("hierarchy_columns", []),
                    "value_columns": parse_plan.get("value_columns", artifact.get("candidate_fields", [])),
                },
                ensure_ascii=False,
                default=str,
            )
        )
        hierarchy_definition = str(artifact.get("hierarchy_definition") or json.dumps(parse_plan, ensure_ascii=False, default=str))
        return {
            "kb_id": artifact["kb_id"],
            "file_id": artifact["file_id"],
            "table_id": artifact["table_id"],
            "table_title": artifact["table_title"],
            "file_name": artifact["file_name"],
            "sheet_name": artifact["sheet_name"],
            "summary_text": artifact["summary_text"],
            "candidate_fields": artifact["candidate_fields"],
            "tree_metric_names": artifact["tree_metric_names"],
            "tree_path_text": artifact["tree_path_text"],
            "tree_leaf_text": tree_fields.get("tree_leaf_text", []),
            "tree_search_text": tree_fields.get("tree_search_text", ""),
            "normalized_headers": normalized_headers,
            "hierarchy_definition": hierarchy_definition,
            "tree_object_name": artifact["tree_object_name"],
            "source_object_name": artifact.get("source_object_name", ""),
            "normalized_object_name": artifact.get("normalized_object_name", ""),
            "tree_object": artifact["tree_object_name"],
            "source_object": artifact.get("source_object_name", ""),
            "xlsx_object": artifact.get("normalized_object_name", ""),
            "parse_mode": str(artifact.get("parse_mode") or _parse_mode(parse_plan)),
            "large_table_reason": artifact.get("large_table_reason", ""),
            "embedding_error": "",
            "row_count": artifact["row_count"],
            "column_count": artifact["column_count"],
            "source": "artifact_fallback",
        }

    def search_tables(
        self,
        db: Session,
        *,
        question: str,
        kb_id: int | None = None,
        top_k: int = 8,
    ) -> list[dict[str, Any]]:
        query = str(question or "").strip()
        if not query:
            return []

        from services.common.knowledge_guard_service import knowledge_guard
        from core.knowledge_policy import KnowledgeOperation
        kb_repo = KnowledgeBaseRepository(db)
        if kb_id:
            guard_ctx = knowledge_guard.resolve_kb_for_operation(db, int(kb_id), KnowledgeOperation.TABLE_READ)
            kbs = [guard_ctx.kb]
        else:
            kbs = kb_repo.list_by_usage(KB_USAGE_TABLE_SEMANTIC_TREE)

        hits: list[dict[str, Any]] = []
        for kb in kbs:
            try:
                hits.extend(
                    table_search_index_service.search(
                        db,
                        collection_name=str(kb.collection_name),
                        kb_id=int(kb.id),
                        question=query,
                        top_k=top_k,
                    )
                )
            except Exception as exc:  # noqa: BLE001
                logger.warning("table search index query failed: kb_id=%s error=%s", kb.id, exc)

        if hits:
            hits.sort(key=lambda item: float(item.get("score") or 0.0), reverse=True)
            return hits[: max(1, min(int(top_k), 50))]

        rows = self.list_tables(db, kb_id=kb_id, limit=200)
        fallback_hits = []
        for row in rows:
            payload = self.artifact_to_response(row)
            text = "\n".join(
                [
                    payload["table_title"],
                    payload["summary_text"],
                    "、".join(payload["candidate_fields"]),
                    "、".join(payload["tree_metric_names"]),
                    "\n".join(payload["tree_path_text"][:80]),
                ]
            )
            score = _lexical_table_score(query, text)
            if score <= 0:
                continue
            fallback_hits.append({**payload, "score": score, "source": "artifact_lexical"})
        fallback_hits.sort(key=lambda item: float(item.get("score") or 0.0), reverse=True)
        return fallback_hits[: max(1, min(int(top_k), 50))]

    def list_jobs(
        self,
        db: Session,
        *,
        limit: int = 50,
    ) -> list[dict[str, Any]]:
        kb_rows = (
            KnowledgeBaseRepository(db)
            .list_by_usage(
                KB_USAGE_TABLE_SEMANTIC_TREE
            )
        )

        kb_ids = [
            int(item.id)
            for item in kb_rows
        ]

        if not kb_ids:
            return []

        safe_limit = max(
            1,
            min(int(limit), 200),
        )

        files = (
            db.query(KnowledgeFile)
            .filter(
                KnowledgeFile.kb_id.in_(kb_ids),
                KnowledgeFile.is_deleted
                == False,  # noqa: E712
                KnowledgeFile.usage_snapshot
                == KB_USAGE_TABLE_SEMANTIC_TREE,
                KnowledgeFile.processor_type
                == ProcessorType.TABLE_SEMANTIC.value,
            )
            .order_by(
                KnowledgeFile.created_at.desc(),
                KnowledgeFile.id.desc(),
            )
            .limit(safe_limit)
            .all()
        )

        return [
            self._file_to_job_response(
                db,
                item,
            )
            for item in files
        ]

    def get_job_by_task_id(
        self,
        db: Session,
        task_id: str,
    ) -> dict[str, Any]:
        file_entity = (
            KnowledgeFileRepository(db)
            .get_by_task_id(task_id)
        )

        if file_entity is None:
            raise ValueError(
                "Table task not found"
            )

        return self._file_to_job_response(
            db,
            file_entity,
        )

    def get_job_by_table_id(self, db: Session, table_id: str) -> dict[str, Any]:
        entity = self.get_table(db, table_id)
        file_entity = KnowledgeFileRepository(db).get_by_id(int(entity.file_id))
        if file_entity is None:
            raise ValueError("Source file not found")
        return self._file_to_job_response(db, file_entity, table_id=str(table_id))

    @staticmethod
    def artifact_to_response(
        entity: TableSemanticArtifact,
    ) -> dict:
        return {
            "id": int(entity.id),
            "kb_id": int(entity.kb_id),
            "file_id": int(entity.file_id),

            "usage_snapshot": str(
                entity.usage_snapshot
            ),
            "processor_type": (
                ProcessorType.TABLE_SEMANTIC.value
            ),

            "table_id": str(entity.table_id),
            "file_name": str(entity.file_name),
            "sheet_name": str(entity.sheet_name),
            "table_title": str(entity.table_title),
            "summary_text": str(entity.summary_text),

            "candidate_fields": _loads_list(
                entity.candidate_fields_json
            ),
            "tree_path_text": _loads_list(
                entity.tree_path_text_json
            ),
            "tree_metric_names": _loads_list(
                entity.tree_metric_names_json
            ),

            "tree_object_name": str(
                entity.tree_object_name
            ),
            "row_count": int(entity.row_count),
            "column_count": int(
                entity.column_count
            ),
            "created_at": entity.created_at,
            "updated_at": entity.updated_at,
        }

    def reprocess_table_file(
        self,
        *,
        db: Session,
        file_id: int,
    ) -> str:
        guard_ctx = (
            knowledge_guard.resolve_file_for_operation(
                db,
                int(file_id),
                KnowledgeOperation.TABLE_REPROCESS,
            )
        )

        file_entity = guard_ctx.file

        if file_entity is None:
            raise ValueError(
                "Table source file not found"
            )

        if (
            guard_ctx.processor_type
            != ProcessorType.TABLE_SEMANTIC
        ):
            raise FileProcessorMismatchError(
                "Table reprocess requires "
                "table_semantic processor"
            )

        artifact_repo = (
            TableSemanticArtifactRepository(db)
        )

        artifacts = (
            artifact_repo.list_valid_by_file_id(
                file_id=int(file_entity.id),
                limit=500,
            )
        )

        # 已存在产物时保留 table_id，避免外部引用失效。
        if artifacts:
            sheet_jobs = [
                {
                    "table_id": str(
                        artifact.table_id
                    ),
                    "sheet_name": str(
                        artifact.sheet_name or ""
                    ).strip() or None,
                }
                for artifact in artifacts
            ]
        else:
            raw = minio_repo.get_file_bytes(
                str(file_entity.minio_object_name)
            )

            sheet_jobs = self._build_sheet_jobs(
                raw=raw,
                file_type=str(
                    file_entity.file_type
                ),
                sheet_name=None,
            )

        celery_task_id = uuid4().hex
        batch_id = uuid4().hex

        file_repo = KnowledgeFileRepository(db)

        process_version = file_repo.submit_task(
            int(file_entity.id),
            celery_task_id,
        )

        table_ids = [
            str(item.get("table_id") or "")
            for item in sheet_jobs
        ]

        sheet_names = [
            str(item.get("sheet_name") or "")
            for item in sheet_jobs
        ]

        initial_task_meta = {
            "file_id": int(file_entity.id),
            "batch_id": batch_id,
            "state": "PENDING",

            "table_ids": table_ids,
            "sheet_names": sheet_names,

            "total_sheets": len(sheet_jobs),
            "completed_sheets": 0,
            "progress": 0.0,

            "completed_table_ids": [],
            "current_table_id": None,
            "current_sheet_name": None,

            "error": None,
        }

        persisted = (
            file_repo
            .update_table_task_meta_for_task(
                int(file_entity.id),
                task_id=celery_task_id,
                process_version=process_version,
                task_meta=initial_task_meta,
            )
        )

        if not persisted:
            raise RuntimeError(
                "无法持久化表格重处理任务计划"
            )

        from tasks.table_semantic_tasks import (
            process_table_semantic_task,
        )

        try:
            process_table_semantic_task.apply_async(
                kwargs={
                    "file_id": int(file_entity.id),
                    "batch_id": batch_id,
                    "sheet_jobs": sheet_jobs,
                    "expected_usage": (
                        KB_USAGE_TABLE_SEMANTIC_TREE
                    ),
                    "expected_processor_type": (
                        ProcessorType
                        .TABLE_SEMANTIC
                        .value
                    ),
                    "process_version": (
                        process_version
                    ),
                },
                task_id=celery_task_id,
            )
        except Exception as exc:
            file_repo.update_status_for_task(
                int(file_entity.id),
                task_id=celery_task_id,
                process_version=process_version,
                status=3,
                error_msg=(
                    "Failed to publish table "
                    f"reprocess task: {exc}"
                )[:1000],
            )
            raise

        return celery_task_id

    @staticmethod
    def _resolve_table_kb(db: Session, kb_id: int | None):
        from services.common.knowledge_guard_service import knowledge_guard
        from core.knowledge_policy import KnowledgeOperation
        from core.knowledge_usage import KB_USAGE_TABLE_SEMANTIC_TREE

        if kb_id is None:
            raise ValueError("上传表格时必须明确指定表格语义树知识库 ID")

        guard_ctx = knowledge_guard.resolve_kb_for_operation(
            db,
            int(kb_id),
            KnowledgeOperation.TABLE_UPLOAD,
        )
        return guard_ctx.kb

    def _build_sheet_jobs(self, *, raw: bytes, file_type: str, sheet_name: str | None) -> list[dict[str, str | None]]:
        sheet_names = self._detect_sheet_names(raw=raw, file_type=file_type)
        if sheet_name:
            if sheet_names and sheet_name not in sheet_names:
                raise ValueError(f"工作表不存在：{sheet_name}")
            sheet_names = [sheet_name]
        if not sheet_names:
            sheet_names = [None]
        return [{"table_id": uuid4().hex, "sheet_name": name} for name in sheet_names]

    @staticmethod
    def _detect_sheet_names(*, raw: bytes, file_type: str) -> list[str | None]:
        ext = str(file_type or "").lower().strip(".")
        if ext in {"xlsx", "xlsm"}:
            workbook = load_workbook(io.BytesIO(raw), data_only=True, read_only=False)
            return [str(name) for name in workbook.sheetnames]
        if ext == "xls":
            try:
                import pandas as pd
            except Exception as exc:  # noqa: BLE001
                raise ValueError("解析 .xls 文件需要 pandas/xlrd 依赖") from exc
            excel = pd.ExcelFile(io.BytesIO(raw))
            return [str(name) for name in excel.sheet_names]
        return [None]

    @staticmethod
    def _build_normalized_workbook(parsed: ParsedTable) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = _safe_excel_sheet_title(parsed.sheet_name)
        sheet.append(parsed.headers)
        for record in parsed.records:
            sheet.append([record.get(header) for header in parsed.headers])
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def _file_to_job_response(
        self,
        db: Session,
        file_entity: KnowledgeFile,
        *,
        table_id: str | None = None,
    ) -> dict[str, Any]:
        artifacts = (
            TableSemanticArtifactRepository(db)
            .list_valid_by_file_id(
                file_id=int(file_entity.id),
                limit=200,
            )
        )

        task_meta = _loads_dict(
            file_entity.table_task_meta_json
        )

        artifact_table_ids = [
            str(item.table_id)
            for item in artifacts
        ]

        table_ids = [
            str(item)
            for item in (
                task_meta.get("table_ids")
                or artifact_table_ids
            )
            if item
        ]

        sheet_names = [
            str(item)
            for item in (
                task_meta.get("sheet_names")
                or [
                    item.sheet_name
                    for item in artifacts
                ]
            )
            if item is not None
        ]

        total_sheets = int(
            task_meta.get("total_sheets")
            if task_meta.get("total_sheets")
            is not None
            else (
                len(table_ids)
                or len(sheet_names)
            )
        )

        completed_sheets = int(
            task_meta.get("completed_sheets")
            if task_meta.get("completed_sheets")
            is not None
            else len(artifacts)
        )

        progress = float(
            task_meta.get("progress")
            if task_meta.get("progress")
            is not None
            else (
                completed_sheets / total_sheets
                if total_sheets
                else 0.0
            )
        )

        return {
            "file_id": int(file_entity.id),
            "task_id": str(
                file_entity.task_id or ""
            ),
            "batch_id": str(
                task_meta.get("batch_id") or ""
            ),

            # 数据库文件状态是持久状态来源。
            "state": _file_status_text(
                int(file_entity.status)
            ),
            "ready": int(
                file_entity.status
            ) in {2, 3},
            "successful": int(
                file_entity.status
            ) == 2,

            "error": (
                file_entity.error_msg
                or task_meta.get("error")
            ),

            "file_name": str(
                file_entity.file_name
            ),

            "table_id": table_id,
            "table_ids": table_ids,
            "sheet_names": sheet_names,

            "total_sheets": total_sheets,
            "completed_sheets": (
                completed_sheets
            ),
            "progress": progress,

            "completed_table_ids": [
                str(item)
                for item in (
                    task_meta.get(
                        "completed_table_ids"
                    )
                    or artifact_table_ids
                )
                if item
            ],

            "current_table_id": (
                task_meta.get(
                    "current_table_id"
                )
            ),
            "current_sheet_name": (
                task_meta.get(
                    "current_sheet_name"
                )
            ),

            "tables": [
                self.artifact_to_response(item)
                for item in artifacts
            ],

            "created_at": file_entity.created_at,
            "updated_at": file_entity.updated_at,
        }

    @staticmethod
    def _load_rows(raw: bytes, *, file_type: str, sheet_name: str | None) -> tuple[list[list[Any]], str]:
        ext = str(file_type or "").lower().strip(".")
        if ext == "csv":
            text = _decode_bytes(raw)
            reader = csv.reader(io.StringIO(text))
            return [list(row) for row in reader], sheet_name or "CSV"

        if ext in {"xlsx", "xlsm"}:
            workbook = load_workbook(io.BytesIO(raw), data_only=True, read_only=False)
            if sheet_name and sheet_name not in workbook.sheetnames:
                raise ValueError(f"工作表不存在：{sheet_name}")
            selected_sheet = sheet_name or workbook.sheetnames[0]
            sheet = workbook[selected_sheet]
            rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
            return rows, selected_sheet

        if ext == "xls":
            try:
                import pandas as pd
            except Exception as exc:  # noqa: BLE001
                raise ValueError("解析 .xls 文件需要 pandas/xlrd 依赖") from exc
            data = pd.read_excel(io.BytesIO(raw), sheet_name=sheet_name or 0, header=None)
            resolved_sheet = sheet_name or "Sheet1"
            return data.where(data.notna(), None).values.tolist(), resolved_sheet

        raise ValueError(f"Unsupported table file type: .{ext}")

    @staticmethod
    def _from_plan_result(*, result: PlanBuildResult, file_name: str) -> ParsedTable:
        parse_plan = _compatible_parse_plan(result)
        summary_text = TableSemanticService._build_summary(
            title=result.title,
            file_name=file_name,
            sheet_name=result.sheet_name,
            row_count=result.row_count,
            headers=result.headers,
        )
        tree = dict(result.tree)
        tree["摘要"] = summary_text
        tree["字段"] = result.headers
        tree_with_cell_refs = dict(result.tree_with_cell_refs)
        tree_with_cell_refs["摘要"] = summary_text
        tree_with_cell_refs["字段"] = result.headers
        tree_with_cell_refs["路径"] = result.tree_path_text[:MAX_TREE_PATHS]
        return ParsedTable(
            sheet_name=result.sheet_name,
            title=result.title,
            summary_text=summary_text,
            headers=result.headers,
            records=result.records,
            tree=tree,
            tree_with_cell_refs=tree_with_cell_refs,
            tree_path_text=result.tree_path_text[:MAX_TREE_PATHS],
            tree_metric_names=result.tree_metric_names,
            row_count=result.row_count,
            column_count=result.column_count,
            parse_plan=parse_plan,
            parse_mode=str(parse_plan.get("source") or "plan_based"),
            coverage={
                **asdict(result.coverage),
                "complete": bool(result.coverage.complete),
                "is_complete": bool(result.coverage.complete),
                "data_rows": int(result.coverage.data_rows or result.row_count),
                "value_columns": int(
                    result.coverage.value_columns or len(result.parse_plan.value_columns)
                ),
                "skipped_rows": list(result.coverage.skipped_rows),
            },
            validation_warnings=list(result.validation_warnings),
        )

    @staticmethod
    def _from_enhanced_result(
        *,
        result: EnhancedTableParseResult,
        sheet,
        file_name: str,
        markdown_table: str,
    ) -> ParsedTable:
        plan_result = table_parse_plan_service.parse_sheet(sheet, file_name=file_name)
        table_title = str(result.table_title or plan_result.title or os.path.splitext(file_name)[0]).strip()
        summary_text = str(result.summary_text or "").strip() or TableSemanticService._build_summary(
            title=table_title,
            file_name=file_name,
            sheet_name=plan_result.sheet_name,
            row_count=plan_result.row_count,
            headers=plan_result.headers,
        )
        headers = _headers_from_normalized(result.normalized_headers) or plan_result.headers
        tree = _attach_enhanced_metadata(
            result.tree,
            table_title=table_title,
            summary_text=summary_text,
            headers=headers,
        )
        tree_with_cell_refs = _attach_enhanced_metadata(
            result.tree_with_cell_refs,
            table_title=table_title,
            summary_text=summary_text,
            headers=headers,
        )
        tree_fields = build_tree_index_fields(tree)
        parse_plan = _compatible_parse_plan(plan_result)
        parse_plan.update(
            {
                "source": "enhanced_llm",
                "normalized_headers": result.normalized_headers,
                "hierarchy_definition": result.hierarchy_definition,
                "final_json_tree": result.final_json_tree,
            }
        )
        tree_path_text = _string_list(tree_fields.get("tree_path_text")) or plan_result.tree_path_text[:MAX_TREE_PATHS]
        tree_metric_names = _string_list(tree_fields.get("tree_metric_names")) or plan_result.tree_metric_names or headers
        return ParsedTable(
            sheet_name=plan_result.sheet_name,
            title=table_title,
            summary_text=summary_text,
            headers=headers,
            records=plan_result.records,
            tree=tree,
            tree_with_cell_refs=tree_with_cell_refs,
            tree_path_text=tree_path_text[:MAX_TREE_PATHS],
            tree_metric_names=tree_metric_names,
            row_count=plan_result.row_count,
            column_count=plan_result.column_count,
            parse_plan=parse_plan,
            parse_mode="enhanced_llm",
            markdown_table=markdown_table or result.markdown_table,
            normalized_headers=result.normalized_headers,
            hierarchy_definition=result.hierarchy_definition,
            final_json_tree=result.final_json_tree,
            coverage={
                **asdict(plan_result.coverage),
                "complete": bool(plan_result.coverage.complete),
                "is_complete": bool(plan_result.coverage.complete),
                "data_rows": int(plan_result.coverage.data_rows or plan_result.row_count),
                "value_columns": int(
                    plan_result.coverage.value_columns or len(plan_result.parse_plan.value_columns)
                ),
                "skipped_rows": list(plan_result.coverage.skipped_rows),
            },
            validation_warnings=list(plan_result.validation_warnings),
        )

    @staticmethod
    def _trim_grid(rows: list[list[Any]]) -> list[list[Any]]:
        non_empty_rows = [list(row) for row in rows if any(not _is_empty(cell) for cell in row)]
        if not non_empty_rows:
            return []

        max_len = max(len(row) for row in non_empty_rows)
        padded = [row + [None] * (max_len - len(row)) for row in non_empty_rows]
        used_cols = [
            index
            for index in range(max_len)
            if any(not _is_empty(row[index]) for row in padded)
        ]
        return [[row[index] for index in used_cols] for row in padded]

    @staticmethod
    def _detect_header(rows: list[list[Any]], *, fallback_title: str) -> tuple[int, str]:
        for index, row in enumerate(rows):
            non_empty = [_cell_text(cell) for cell in row if not _is_empty(cell)]
            if len(non_empty) >= 2:
                title = ""
                if index > 0:
                    previous = [_cell_text(cell) for cell in rows[index - 1] if not _is_empty(cell)]
                    if 0 < len(previous) <= 2:
                        title = " ".join(previous)
                return index, title or os.path.splitext(fallback_title)[0]
        return 0, os.path.splitext(fallback_title)[0]

    @staticmethod
    def _normalize_headers(row: list[Any]) -> list[str]:
        headers: list[str] = []
        seen: dict[str, int] = {}
        for index, cell in enumerate(row):
            base = _cell_text(cell) or f"列{index + 1}"
            count = seen.get(base, 0)
            seen[base] = count + 1
            headers.append(base if count == 0 else f"{base}_{count + 1}")
        return headers

    @staticmethod
    def _build_records(headers: list[str], rows: list[list[Any]]) -> list[dict[str, Any]]:
        records: list[dict[str, Any]] = []
        for row in rows:
            if all(_is_empty(cell) for cell in row):
                continue
            record: dict[str, Any] = {}
            for index, header in enumerate(headers):
                value = row[index] if index < len(row) else None
                record[header] = _safe_json(value)
            records.append(record)
        return records

    @staticmethod
    def _build_summary(
        *,
        title: str,
        file_name: str,
        sheet_name: str,
        row_count: int,
        headers: list[str],
    ) -> str:
        fields = "、".join(headers[:80])
        return (
            f"表格《{title}》来自文件 {file_name} 的工作表 {sheet_name}，"
            f"共 {row_count} 行数据、{len(headers)} 个字段。字段包括：{fields}。"
        )

    @staticmethod
    def _build_tree_paths(*, title: str, records: list[dict[str, Any]], limit: int) -> list[str]:
        paths: list[str] = []
        for row_index, record in enumerate(records, start=1):
            row_label = _choose_row_label(record, row_index)
            for field, value in record.items():
                if _is_empty(value):
                    continue
                paths.append(f"{title} | {row_label} | {field}: {_cell_text(value)}")
                if len(paths) >= limit:
                    return paths
        return paths

    @staticmethod
    def _build_index_chunks(table_id: str, parsed: ParsedTable) -> list[str]:
        chunks = [
            (
                f"table_id: {table_id}\n"
                f"表格标题: {parsed.title}\n"
                f"工作表: {parsed.sheet_name}\n"
                f"表格摘要: {parsed.summary_text}"
            ),
            (
                f"table_id: {table_id}\n"
                f"表格标题: {parsed.title}\n"
                f"字段: {'、'.join(parsed.headers)}"
            ),
        ]
        for start in range(0, len(parsed.tree_path_text), PATHS_PER_CHUNK):
            piece = parsed.tree_path_text[start : start + PATHS_PER_CHUNK]
            chunks.append(
                f"table_id: {table_id}\n表格标题: {parsed.title}\n树路径:\n"
                + "\n".join(piece)
            )
        return chunks

    @staticmethod
    def _index_chunks(db: Session, collection_name: str, chunks: list[DocumentChunk]) -> None:
        if not chunks:
            return
        embeddings = get_embeddings(db).embed_documents([item.content for item in chunks])
        if len(embeddings) != len(chunks):
            raise ValueError("Embedding generation returned unexpected vector count")
        vector_dim = int(get_embedding_vector_dim(db))
        rows = [
            {
                "chunk_id": int(chunk.id),
                "kb_id": int(chunk.kb_id),
                "file_id": int(chunk.file_id),
                "text": chunk.content,
                "embedding": vector,
                "usage_type": chunk.usage_snapshot,
                "processor_type": chunk.processor_type,
            }
            for chunk, vector in zip(chunks, embeddings)
        ]
        es_repo.insert_chunks(rows, index_name=collection_name, vector_dim=vector_dim)


def _decode_bytes(raw: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "gbk"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def _attach_enhanced_metadata(
    tree: dict[str, Any],
    *,
    table_title: str,
    summary_text: str,
    headers: list[str],
) -> dict[str, Any]:
    payload = dict(tree) if isinstance(tree, dict) else {"data": tree}
    payload.setdefault("\u8868\u683c\u540d\u79f0", table_title)
    payload.setdefault("\u6458\u8981", summary_text)
    payload.setdefault("\u5b57\u6bb5", headers)
    return payload


def _headers_from_normalized(value: str) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        return []
    headers: list[str] = []

    def collect(item: Any) -> None:
        if isinstance(item, str):
            headers.append(item)
            return
        if isinstance(item, dict):
            raw = item.get("header") or item.get("name") or item.get("field") or item.get("title")
            if raw:
                headers.append(str(raw))

    if isinstance(parsed, list):
        for item in parsed:
            collect(item)
    elif isinstance(parsed, dict):
        for key in ("hierarchy_keys", "value_leaves", "hierarchy_columns", "value_columns"):
            values = parsed.get(key)
            if isinstance(values, list):
                for item in values:
                    collect(item)
    return _dedupe_ordered(headers)


def _string_list(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(item) for item in value if str(item).strip()]


def _dedupe_ordered(values: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        text = str(value or "").strip()
        if not text or text in seen:
            continue
        seen.add(text)
        result.append(text)
    return result


def _choose_row_label(record: dict[str, Any], row_index: int) -> str:
    for key, value in record.items():
        if not _is_empty(value) and not re.search(r"金额|数量|总数|比例|率|分数|得分|面积|均值|平均", str(key)):
            return f"第{row_index}行（{key}={_cell_text(value)}）"
    return f"第{row_index}行"


def _safe_excel_sheet_title(value: str) -> str:
    text = re.sub(r"[\[\]:*?/\\]", "_", str(value or "Sheet1").strip())
    return (text or "Sheet1")[:31]


def _file_status_text(status: int) -> str:
    return {
        0: "PENDING",
        1: "STARTED",
        2: "SUCCESS",
        3: "FAILURE",
    }.get(int(status), "UNKNOWN")


def _lexical_table_score(query: str, text: str) -> float:
    q = str(query or "").lower()
    t = str(text or "").lower()
    if not q or not t:
        return 0.0
    terms = [term for term in re.split(r"[\s,，。；;：:、|/\\()（）]+", q) if term]
    score = sum(2.0 for term in terms if len(term) > 1 and term in t)
    query_chars = {char for char in q if "\u4e00" <= char <= "\u9fff"}
    if query_chars:
        text_chars = {char for char in t if "\u4e00" <= char <= "\u9fff"}
        score += len(query_chars & text_chars) / max(1, len(query_chars))
    return float(score)


def _parse_mode(parse_plan: dict[str, Any] | None) -> str:
    if not isinstance(parse_plan, dict) or not parse_plan:
        return "heuristic"
    source = str(parse_plan.get("source") or "").strip()
    if source:
        return source
    return "plan_based"


def _compatible_parse_plan(result: PlanBuildResult) -> dict[str, Any]:
    plan = result.parse_plan
    if hasattr(plan, "model_dump"):
        payload = plan.model_dump(mode="json")
        payload.update(
            {
                "source": "llm",
                "raw_plan_output": str(result.raw_plan_output or "")[:8000],
            }
        )
        return payload

    columns = [*plan.hierarchy_columns, *plan.value_columns]
    min_col = min(columns) if columns else 1
    max_col = min_col + max(0, int(result.column_count) - 1)
    max_row = max(plan.data_start_row, plan.data_start_row + max(0, int(result.row_count) - 1))
    header_range = f"{get_column_letter(min_col)}{plan.header_row}:{get_column_letter(max_col)}{plan.header_row}"
    data_row_range = f"{get_column_letter(min_col)}{plan.data_start_row}:{get_column_letter(max_col)}{max_row}"
    title_ranges = []
    if plan.header_row > 1:
        title_ranges.append(f"{get_column_letter(min_col)}1:{get_column_letter(max_col)}{plan.header_row - 1}")
    return {
        "title": plan.title,
        "header_row": plan.header_row,
        "data_start_row": plan.data_start_row,
        "hierarchy_columns_raw": list(plan.hierarchy_columns),
        "value_columns_raw": list(plan.value_columns),
        "fill_down": bool(plan.fill_down),
        "source": plan.source,
        "raw_plan_output": plan.raw_plan_output,
        "table_range": f"{get_column_letter(min_col)}1:{get_column_letter(max_col)}{max_row}",
        "title_ranges": title_ranges,
        "header_ranges": [header_range],
        "data_row_range": data_row_range,
        "hierarchy_columns": [
            {
                "header": _header_for_result(result, col_idx, min_col),
                "col": get_column_letter(col_idx),
            }
            for col_idx in plan.hierarchy_columns
        ],
        "value_columns": [
            {
                "header": _header_for_result(result, col_idx, min_col),
                "col": get_column_letter(col_idx),
                "group": _header_for_result(result, col_idx, min_col),
            }
            for col_idx in plan.value_columns
        ],
        "row_paths": [],
        "ignored_rows": [],
        "hierarchy_fill_down": bool(plan.fill_down),
        "notes": list(result.validation_warnings),
    }


def _header_for_result(result: PlanBuildResult, col_idx: int, min_col: int) -> str:
    offset = int(col_idx) - int(min_col)
    if 0 <= offset < len(result.headers):
        return str(result.headers[offset])
    return f"列{offset + 1}"


def _loads_dict(
    value: str | None,
) -> dict[str, Any]:
    if not value:
        return {}

    try:
        payload = json.loads(value)
    except (
        TypeError,
        ValueError,
        json.JSONDecodeError,
    ):
        return {}

    return (
        payload
        if isinstance(payload, dict)
        else {}
    )


def _loads_list(value: str | None) -> list[str]:
    try:
        parsed = json.loads(value or "[]")
    except json.JSONDecodeError:
        return []
    if not isinstance(parsed, list):
        return []
    return [str(item) for item in parsed]
