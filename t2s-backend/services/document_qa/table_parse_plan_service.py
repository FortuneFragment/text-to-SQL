from __future__ import annotations

import json
import os
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from json import JSONDecodeError
from typing import Any

from langchain_core.messages import BaseMessage, HumanMessage
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, Field, ValidationError, field_validator, model_validator

from services.common.llm_service import common_llm_service


TABLE_PARSE_PLAN_PROMPT = """
你是一名电子表格结构解析专家。
你会收到一个 Excel 工作表的 compact profile。profile 可能只包含采样行，但后续程序会在完整工作表上执行解析计划。
你的任务不是生成最终语义树，也不是逐格填写数据。你的任务是输出一个 TableParsePlan，让 Python 根据坐标规则扫描完整 sheet 并确定性构建树。

请严格遵守：
1. 只输出一个合法 JSON 对象，不要输出 Markdown 代码块、解释文字或注释。
2. 表头、行头、分组名称必须保留原始文字，不要翻译、概括、改写或创造新名称。
3. data_row_range 必须覆盖完整工作表中的真实数据行范围，不只覆盖采样行。
4. value_columns 必须包含所有值列，即使该列在采样行中全为空也要包含。
5. hierarchy_columns 表示用于构建完整行路径的列，例如名称、代码、类别、地区、年龄、ID 等。多列行头必须按从左到右包含。
6. 大表优先不要输出逐行 row_paths。只要可以用 hierarchy_columns + value_columns 规则覆盖完整数据区，就让 row_paths 为空。
7. 只有当某些采样行无法用列规则表达父子关系时，才为这些特殊行输出 row_paths；不要为每一行生成 row_paths。
8. ignored_rows 用于列出 data_row_range 内需要跳过的标题行、单位行、注释行、空白分隔行、小节标题行。不要把这些行当数据写树。
9. 如果层级列中的空白单元格表示沿用上一行或上一合并区域的值，请将 hierarchy_fill_down 设置为 true。
10. 如果输入 profile 显示一个 sheet 内有多个形状不同的区域，请选择主数据区域作为 table_range/data_row_range，并在 notes 说明其它区域需要另行切分；不要强行用一个 plan 覆盖多个异构区域。

输出 JSON 结构必须符合以下格式：
{
  "table_range": "A1:T21",
  "title_ranges": ["A1:T2"],
  "header_ranges": ["A3:T4"],
  "data_row_range": "A6:T21",
  "hierarchy_columns": [
    {"header": "指标名称", "col": "A"},
    {"header": "代码", "col": "B"}
  ],
  "value_columns": [
    {"header": "高职专科", "col": "C", "group": "招生数"},
    {"header": "高职专科 - #女", "col": "D", "group": "招生数"}
  ],
  "row_paths": [],
  "ignored_rows": [],
  "hierarchy_fill_down": true,
  "notes": []
}

注意：
- row_paths 可以为空。大表默认应为空，除非有少量特殊行需要手工路径。
- row_paths 中的 ref 必须是单个 Excel 单元格坐标，例如 "A10"。
- value_columns 中的 col 必须是 Excel 列字母，例如 "C"。
- ignored_rows 必须是行号数组，例如 [5, 9]。
- 不要在 JSON 中使用尾逗号。
- 不要输出 JSON 之外的任何内容。

[工作表 compact profile]
{worksheet_json}
"""


@dataclass
class TableParsePlan:
    title: str
    header_row: int
    data_start_row: int
    hierarchy_columns: list[int]
    value_columns: list[int]
    fill_down: bool = True
    source: str = "heuristic"
    raw_plan_output: str = ""


class SheetCell(BaseModel):
    coord: str
    row: int
    col: int
    col_letter: str
    value: Any = None
    effective_value: Any = None
    value_type: str
    merged_range: str | None = None
    is_merged_child: bool = False
    is_blank: bool = False
    number_format: str | None = None
    font_bold: bool = False
    horizontal_alignment: str | None = None


class SheetGrid(BaseModel):
    sheet_name: str
    max_row: int
    max_column: int
    table_range: str
    merged_ranges: list[str]
    rows: list[list[SheetCell]]


class HierarchyColumn(BaseModel):
    header: str
    col: str

    @field_validator("col")
    @classmethod
    def normalize_col(cls, value: str) -> str:
        normalized = value.strip().upper()
        if not re.fullmatch(r"[A-Z]+", normalized):
            raise ValueError("Column must be an Excel column letter")
        return normalized


class ValueColumn(BaseModel):
    header: str
    col: str
    group: str | None = None

    @field_validator("col")
    @classmethod
    def normalize_col(cls, value: str) -> str:
        normalized = value.strip().upper()
        if not re.fullmatch(r"[A-Z]+", normalized):
            raise ValueError("Column must be an Excel column letter")
        return normalized


class RowPathSegment(BaseModel):
    header: str
    ref: str | None = None
    value: Any = None
    role: str | None = None

    @field_validator("ref")
    @classmethod
    def normalize_ref(cls, value: str | None) -> str | None:
        if value is None:
            return None
        normalized = value.strip().upper()
        if not re.fullmatch(r"\$?[A-Z]+\$?[0-9]+", normalized):
            raise ValueError("ref must be a single Excel cell reference")
        return normalized

    @model_validator(mode="after")
    def require_ref_or_value(self) -> RowPathSegment:
        if self.ref is None and self.value is None:
            raise ValueError("Row path segment requires either ref or value")
        return self


class RowPath(BaseModel):
    row: int
    path: list[RowPathSegment]

    @model_validator(mode="after")
    def require_path(self) -> RowPath:
        if not self.path:
            raise ValueError("row path cannot be empty")
        return self


class ComplexTableParsePlan(BaseModel):
    table_range: str
    title_ranges: list[str] = Field(default_factory=list)
    header_ranges: list[str] = Field(default_factory=list)
    data_row_range: str
    hierarchy_columns: list[HierarchyColumn] = Field(default_factory=list)
    value_columns: list[ValueColumn]
    row_paths: list[RowPath] = Field(default_factory=list)
    ignored_rows: list[int] = Field(default_factory=list)
    hierarchy_fill_down: bool = True
    notes: list[str] = Field(default_factory=list)

    @model_validator(mode="after")
    def require_columns(self) -> ComplexTableParsePlan:
        if not self.hierarchy_columns and not self.row_paths:
            raise ValueError("Either hierarchy_columns or row_paths must be provided")
        if not self.value_columns:
            raise ValueError("value_columns cannot be empty")
        return self


@dataclass
class CoverageReport:
    data_rows: int = 0
    value_columns: int = 0
    expected_cells: int = 0
    covered_cells: int = 0
    missing_cells: list[str] = field(default_factory=list)
    skipped_rows: list[int] = field(default_factory=list)

    @property
    def complete(self) -> bool:
        return self.expected_cells == self.covered_cells and not self.missing_cells


@dataclass
class PlanBuildResult:
    title: str
    sheet_name: str
    headers: list[str]
    records: list[dict[str, Any]]
    tree: dict[str, Any]
    tree_with_cell_refs: dict[str, Any]
    tree_path_text: list[str]
    tree_metric_names: list[str]
    row_count: int
    column_count: int
    parse_plan: TableParsePlan | ComplexTableParsePlan
    coverage: CoverageReport
    validation_warnings: list[str] = field(default_factory=list)
    raw_plan_output: str = ""


class TableParsePlanService:
    def parse_sheet(self, sheet: Worksheet, *, file_name: str) -> PlanBuildResult:
        min_row, max_row, min_col, max_col = self._sheet_bounds(sheet)
        if max_row < min_row or max_col < min_col:
            raise ValueError("表格内容为空")

        header_row = self._detect_header_row(sheet, min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col)
        headers = self._headers(sheet, header_row=header_row, min_col=min_col, max_col=max_col)
        title = self._detect_title(sheet, header_row=header_row, min_col=min_col, max_col=max_col, fallback=file_name)
        data_start_row = header_row + 1
        hierarchy_columns, value_columns = self._classify_columns(
            sheet,
            data_start_row=data_start_row,
            max_row=max_row,
            min_col=min_col,
            max_col=max_col,
        )
        if not value_columns:
            value_columns = [col for col in range(min_col, max_col + 1) if col not in hierarchy_columns]
        if not hierarchy_columns:
            first_value = value_columns[0] if value_columns else min_col
            hierarchy_columns = [first_value]
            value_columns = [col for col in value_columns if col != first_value]

        plan = TableParsePlan(
            title=title,
            header_row=header_row,
            data_start_row=data_start_row,
            hierarchy_columns=hierarchy_columns,
            value_columns=value_columns,
            fill_down=True,
        )
        return self.build_result_from_plan(sheet, file_name=file_name, plan=plan)

    def parse_sheet_with_llm(self, sheet: Worksheet, *, file_name: str, db=None) -> PlanBuildResult:
        model = common_llm_service.get_chat_model(db)
        if model is None:
            raise RuntimeError("未配置 LLM，禁止回退启发式解析计划")

        try:
            grid = extract_sheet_grid(sheet)
            prompt_text = TABLE_PARSE_PLAN_PROMPT.replace("{worksheet_json}", grid_to_llm_json(grid))
            raw_output = _message_to_text(model.invoke([HumanMessage(content=prompt_text)]))
            plan = parse_complex_table_parse_plan(raw_output)
            plan = normalize_complex_parse_plan(plan, sheet)
            warnings = validate_complex_parse_plan(plan, sheet)
            result = build_complex_plan_result(
                sheet,
                file_name=file_name,
                plan=plan,
                raw_plan_output=raw_output,
                validation_warnings=warnings,
            )
            if not result.coverage.complete:
                raise ValueError(
                    "LLM 解析计划未完整覆盖值单元格："
                    f"covered={result.coverage.covered_cells}/"
                    f"expected={result.coverage.expected_cells}, "
                    f"missing={result.coverage.missing_cells[:20]}"
                )
            result.validation_warnings.append("已使用 LLM 解析计划")
            return result
        except Exception as exc:  # noqa: BLE001
            raise RuntimeError(
                "LLM 解析计划失败，禁止回退启发式解析："
                f"{exc.__class__.__name__}: {exc}"
            ) from exc

    def build_result_from_plan(self, sheet: Worksheet, *, file_name: str, plan: TableParsePlan) -> PlanBuildResult:
        min_row, max_row, min_col, max_col = self._sheet_bounds(sheet)
        if max_row < min_row or max_col < min_col:
            raise ValueError("表格内容为空")
        headers = self._headers(sheet, header_row=plan.header_row, min_col=min_col, max_col=max_col)
        records, tree, tree_refs, paths, coverage = self._build_tree_from_plan(
            sheet,
            plan=plan,
            headers=headers,
            min_col=min_col,
            max_row=max_row,
        )
        warnings = self._validate(plan=plan, headers=headers, coverage=coverage)
        return PlanBuildResult(
            title=plan.title,
            sheet_name=str(sheet.title),
            headers=headers,
            records=records,
            tree=tree,
            tree_with_cell_refs=tree_refs,
            tree_path_text=paths,
            tree_metric_names=[headers[col - min_col] for col in plan.value_columns if 0 <= col - min_col < len(headers)],
            row_count=len(records),
            column_count=len(headers),
            parse_plan=plan,
            coverage=coverage,
            validation_warnings=warnings,
        )

    @staticmethod
    def _sheet_bounds(sheet: Worksheet) -> tuple[int, int, int, int]:
        min_row = sheet.max_row
        max_row = 1
        min_col = sheet.max_column
        max_col = 1
        found = False
        for row in sheet.iter_rows():
            for cell in row:
                if _is_blank(cell.value):
                    continue
                found = True
                min_row = min(min_row, int(cell.row))
                max_row = max(max_row, int(cell.row))
                min_col = min(min_col, int(cell.column))
                max_col = max(max_col, int(cell.column))
        if not found:
            return 1, 0, 1, 0
        return min_row, max_row, min_col, max_col

    @staticmethod
    def _detect_header_row(sheet: Worksheet, *, min_row: int, max_row: int, min_col: int, max_col: int) -> int:
        best_row = min_row
        best_score = -1.0
        upper = min(max_row, min_row + 12)
        for row_idx in range(min_row, upper + 1):
            values = [_display(sheet.cell(row_idx, col_idx).value) for col_idx in range(min_col, max_col + 1)]
            non_blank = [value for value in values if value]
            if len(non_blank) < 2:
                continue
            numeric_ratio = sum(1 for value in non_blank if _looks_numeric(value)) / max(1, len(non_blank))
            unique_ratio = len(set(non_blank)) / max(1, len(non_blank))
            score = len(non_blank) * 2.0 + unique_ratio - numeric_ratio * 3.0
            if score > best_score:
                best_score = score
                best_row = row_idx
        return best_row

    @staticmethod
    def _headers(sheet: Worksheet, *, header_row: int, min_col: int, max_col: int) -> list[str]:
        headers: list[str] = []
        seen: dict[str, int] = {}
        for col_idx in range(min_col, max_col + 1):
            base = _display(sheet.cell(header_row, col_idx).value) or f"列{col_idx - min_col + 1}"
            count = seen.get(base, 0)
            seen[base] = count + 1
            headers.append(base if count == 0 else f"{base}_{count + 1}")
        return headers

    @staticmethod
    def _detect_title(sheet: Worksheet, *, header_row: int, min_col: int, max_col: int, fallback: str) -> str:
        for row_idx in range(header_row - 1, max(0, header_row - 5), -1):
            values = [_display(sheet.cell(row_idx, col_idx).value) for col_idx in range(min_col, max_col + 1)]
            non_blank = [value for value in values if value]
            if 0 < len(non_blank) <= 3:
                return " ".join(non_blank)
        return os.path.splitext(os.path.basename(fallback))[0] or "未命名表格"

    @staticmethod
    def _classify_columns(
        sheet: Worksheet,
        *,
        data_start_row: int,
        max_row: int,
        min_col: int,
        max_col: int,
    ) -> tuple[list[int], list[int]]:
        profiles: list[tuple[int, float, float, float]] = []
        sample_end = min(max_row, data_start_row + 80)
        for col_idx in range(min_col, max_col + 1):
            values = [_display(sheet.cell(row_idx, col_idx).value) for row_idx in range(data_start_row, sample_end + 1)]
            non_blank = [value for value in values if value]
            fill_ratio = len(non_blank) / max(1, len(values))
            numeric_ratio = sum(1 for value in non_blank if _looks_numeric(value)) / max(1, len(non_blank))
            avg_len = sum(len(value) for value in non_blank) / max(1, len(non_blank))
            profiles.append((col_idx, fill_ratio, numeric_ratio, avg_len))

        hierarchy: list[int] = []
        values: list[int] = []
        for col_idx, fill_ratio, numeric_ratio, avg_len in profiles:
            if fill_ratio <= 0:
                continue
            if numeric_ratio < 0.45 and len(values) == 0 and len(hierarchy) < 4:
                hierarchy.append(col_idx)
                continue
            if numeric_ratio < 0.2 and avg_len <= 40 and len(hierarchy) < 2 and len(values) == 0:
                hierarchy.append(col_idx)
                continue
            values.append(col_idx)

        if not values and len(hierarchy) > 1:
            values = hierarchy[1:]
            hierarchy = hierarchy[:1]
        return hierarchy, values

    def _build_tree_from_plan(
        self,
        sheet: Worksheet,
        *,
        plan: TableParsePlan,
        headers: list[str],
        min_col: int,
        max_row: int,
    ) -> tuple[list[dict[str, Any]], dict[str, Any], dict[str, Any], list[str], CoverageReport]:
        root: dict[str, Any] = {
            "表格名": plan.title,
            "层级字段": [self._header(headers, min_col, col) for col in plan.hierarchy_columns],
            "指标字段": [self._header(headers, min_col, col) for col in plan.value_columns],
            "数据": {},
        }
        root_refs: dict[str, Any] = {
            "表格名": plan.title,
            "层级字段": root["层级字段"],
            "指标字段": root["指标字段"],
            "数据": {},
        }
        paths: list[str] = []
        records: list[dict[str, Any]] = []
        last_hierarchy: dict[int, str] = {}
        coverage = CoverageReport()

        for row_idx in range(plan.data_start_row, max_row + 1):
            row_values = [_display(sheet.cell(row_idx, col_idx).value) for col_idx in [*plan.hierarchy_columns, *plan.value_columns]]
            if not any(row_values):
                continue

            path_parts: list[str] = []
            for col_idx in plan.hierarchy_columns:
                raw = _display(self._effective_cell_value(sheet, row_idx, col_idx))
                if not raw and plan.fill_down:
                    raw = last_hierarchy.get(col_idx, "")
                if raw:
                    last_hierarchy[col_idx] = raw
                header = self._header(headers, min_col, col_idx)
                path_parts.append(f"{header} - {raw or f'第{row_idx}行'}")
            if not path_parts:
                path_parts = [f"第{row_idx}行"]

            record: dict[str, Any] = {}
            for col_idx in range(min_col, min_col + len(headers)):
                header = self._header(headers, min_col, col_idx)
                record[header] = _json_safe(self._effective_cell_value(sheet, row_idx, col_idx))
            records.append(record)

            leaf: dict[str, Any] = {}
            leaf_refs: dict[str, Any] = {}
            for col_idx in plan.value_columns:
                header = self._header(headers, min_col, col_idx)
                value = self._effective_cell_value(sheet, row_idx, col_idx)
                if _is_blank(value):
                    if len(coverage.missing_cells) < 200:
                        coverage.missing_cells.append(f"{get_column_letter(col_idx)}{row_idx}")
                    continue
                coverage.covered_cells += 1
                leaf[header] = _json_safe(value)
                leaf_refs[header] = f"{get_column_letter(col_idx)}{row_idx}"
                paths.append(f"{plan.title} | {' | '.join(path_parts)} | {header}: {_display(value)}")
            coverage.expected_cells += len(plan.value_columns)

            if leaf:
                _insert_path(root["数据"], path_parts, leaf)
                _insert_path(root_refs["数据"], path_parts, leaf_refs)

        return records, root, root_refs, paths, coverage

    @staticmethod
    def _effective_cell_value(sheet: Worksheet, row_idx: int, col_idx: int) -> Any:
        cell = sheet.cell(row_idx, col_idx)
        if cell.value is not None:
            return cell.value
        for merged_range in sheet.merged_cells.ranges:
            if cell.coordinate in merged_range:
                return sheet.cell(merged_range.min_row, merged_range.min_col).value
        return None

    @staticmethod
    def _header(headers: list[str], min_col: int, col_idx: int) -> str:
        offset = col_idx - min_col
        if 0 <= offset < len(headers):
            return headers[offset]
        return f"列{offset + 1}"

    @staticmethod
    def _validate(*, plan: TableParsePlan, headers: list[str], coverage: CoverageReport) -> list[str]:
        warnings: list[str] = []
        if not plan.hierarchy_columns:
            warnings.append("未识别到层级列")
        if not plan.value_columns:
            warnings.append("未识别到指标列")
        if coverage.expected_cells > 0 and not coverage.complete:
            warnings.append(
                f"值单元格覆盖不完整：expected={coverage.expected_cells}, covered={coverage.covered_cells}"
            )
        if len(set(headers)) != len(headers):
            warnings.append("字段名存在重复")
        return warnings

def _insert_path(root: dict[str, Any], path_parts: list[str], leaf: dict[str, Any]) -> None:
    current = root
    for part in path_parts:
        current = current.setdefault(part, {})
    current.update(leaf)


def _display(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value).strip()


def _is_blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def _looks_numeric(value: Any) -> bool:
    text = _display(value)
    if not text:
        return False
    return bool(re.fullmatch(r"[-+]?\d+(?:\.\d+)?%?", text.replace(",", "")))


def _extract_json_object(text: str) -> str:
    value = str(text or "").strip()
    if value.startswith("```"):
        value = re.sub(r"^```(?:json)?", "", value, flags=re.IGNORECASE).strip()
        value = re.sub(r"```$", "", value).strip()
    start = value.find("{")
    end = value.rfind("}")
    if start < 0 or end < start:
        raise ValueError("未从 LLM 输出中找到 JSON 对象")
    return value[start : end + 1]


def _json_safe(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, list):
        return [_json_safe(item) for item in value]
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def extract_sheet_grid(sheet: Worksheet) -> SheetGrid:
    merged_lookup = _build_merged_lookup(sheet)
    rows: list[list[SheetCell]] = []

    for row_idx in range(1, sheet.max_row + 1):
        row_cells: list[SheetCell] = []
        for col_idx in range(1, sheet.max_column + 1):
            cell = sheet.cell(row_idx, col_idx)
            coord = f"{get_column_letter(col_idx)}{row_idx}"
            merged_info = merged_lookup.get((row_idx, col_idx))
            merged_range = merged_info["range"] if merged_info else None
            is_merged_child = bool(merged_info and not merged_info["is_top_left"])
            effective_value = merged_info["top_left_value"] if merged_info else cell.value
            row_cells.append(
                SheetCell(
                    coord=coord,
                    row=row_idx,
                    col=col_idx,
                    col_letter=get_column_letter(col_idx),
                    value=_json_safe(cell.value),
                    effective_value=_json_safe(effective_value),
                    value_type=type(effective_value).__name__,
                    merged_range=merged_range,
                    is_merged_child=is_merged_child,
                    is_blank=_is_blank(effective_value),
                    number_format=cell.number_format,
                    font_bold=bool(cell.font and cell.font.bold),
                    horizontal_alignment=cell.alignment.horizontal,
                )
            )
        rows.append(row_cells)

    max_column = max(1, int(sheet.max_column or 1))
    max_row = max(1, int(sheet.max_row or 1))
    return SheetGrid(
        sheet_name=str(sheet.title),
        max_row=max_row,
        max_column=max_column,
        table_range=f"A1:{get_column_letter(max_column)}{max_row}",
        merged_ranges=[str(merged_range) for merged_range in sheet.merged_cells.ranges],
        rows=rows,
    )


def grid_to_llm_json(grid: SheetGrid) -> str:
    sample_row_indexes = _sample_row_indexes(grid)
    payload: dict[str, Any] = {
        "sheet_name": grid.sheet_name,
        "max_row": grid.max_row,
        "max_column": grid.max_column,
        "table_range": grid.table_range,
        "merged_ranges": grid.merged_ranges,
        "sampled_rows": [
            _compact_row(grid.rows[row_idx - 1])
            for row_idx in sample_row_indexes
            if 1 <= row_idx <= len(grid.rows)
        ],
        "row_profiles": [_row_profile(row) for row in grid.rows],
        "column_profiles": _column_profiles(grid),
    }
    return json.dumps(payload, ensure_ascii=False)


def _sample_row_indexes(grid: SheetGrid) -> list[int]:
    row_count = grid.max_row
    selected: set[int] = set(range(1, min(row_count, 20) + 1))
    selected.update(range(max(1, row_count - 8), row_count + 1))
    if row_count > 28:
        step = max(1, row_count // 12)
        for row_idx in range(21, row_count - 8, step):
            selected.add(row_idx)
            if row_idx + 1 <= row_count:
                selected.add(row_idx + 1)
    for row in grid.rows:
        if row and any(cell.merged_range and not cell.is_merged_child for cell in row):
            selected.add(row[0].row)
    return sorted(row_idx for row_idx in selected if 1 <= row_idx <= row_count)


def _compact_row(row: list[SheetCell]) -> dict[str, Any]:
    cells: list[dict[str, Any]] = []
    for cell in row:
        if cell.is_merged_child:
            continue
        if cell.is_blank and not cell.merged_range:
            continue
        item: dict[str, Any] = {"c": cell.coord, "v": _compact_value(cell.effective_value)}
        if cell.merged_range:
            item["m"] = cell.merged_range
        if cell.font_bold:
            item["b"] = True
        if cell.horizontal_alignment:
            item["a"] = cell.horizontal_alignment
        cells.append(item)
    return {"row": row[0].row if row else None, "cells": cells}


def _row_profile(row: list[SheetCell]) -> dict[str, Any]:
    non_blank = [cell for cell in row if not cell.is_blank]
    numeric_count = sum(
        1 for cell in non_blank if isinstance(cell.effective_value, (int, float, Decimal))
    )
    return {
        "row": row[0].row if row else None,
        "non_blank": len(non_blank),
        "numeric": numeric_count,
        "text": len(non_blank) - numeric_count,
        "first_values": [
            {"c": cell.coord, "v": _compact_value(cell.effective_value)}
            for cell in non_blank[:4]
        ],
    }


def _column_profiles(grid: SheetGrid) -> list[dict[str, Any]]:
    profiles: list[dict[str, Any]] = []
    for col_idx in range(1, grid.max_column + 1):
        cells = [row[col_idx - 1] for row in grid.rows if len(row) >= col_idx]
        non_blank = [cell for cell in cells if not cell.is_blank]
        numeric_count = sum(
            1 for cell in non_blank if isinstance(cell.effective_value, (int, float, Decimal))
        )
        profiles.append(
            {
                "col": get_column_letter(col_idx),
                "non_blank": len(non_blank),
                "numeric": numeric_count,
                "text": len(non_blank) - numeric_count,
                "first_values": [
                    {"c": cell.coord, "v": _compact_value(cell.effective_value)}
                    for cell in non_blank[:5]
                ],
            }
        )
    return profiles


def _compact_value(value: Any, max_length: int = 80) -> Any:
    if isinstance(value, str):
        text = re.sub(r"\s+", " ", value).strip()
        return text[:max_length] + "..." if len(text) > max_length else text
    return _json_safe(value)


def parse_complex_table_parse_plan(llm_output: str) -> ComplexTableParsePlan:
    json_content = _extract_json_object(llm_output)
    try:
        parsed = json.loads(json_content)
    except JSONDecodeError:
        parsed = json.loads(_repair_common_json_issues(json_content))
    try:
        return ComplexTableParsePlan.model_validate(parsed)
    except ValidationError as exc:
        raise ValueError(f"LLM 解析计划结构校验失败：{exc}") from exc


def normalize_complex_parse_plan(
    plan: ComplexTableParsePlan,
    sheet: Worksheet,
) -> ComplexTableParsePlan:
    row_paths = _repair_row_paths(plan.row_paths, sheet) if plan.row_paths else []
    if not plan.value_columns:
        return plan.model_copy(update={"row_paths": row_paths})

    _, data_min_row, _, data_max_row = range_boundaries(plan.data_row_range)
    table_min_col, _, _, _ = range_boundaries(plan.table_range)
    first_value_col = min(column_index_from_string(column.col) for column in plan.value_columns)
    hierarchy_col_indexes = {
        column_index_from_string(column.col) for column in plan.hierarchy_columns
    }
    value_col_indexes = {column_index_from_string(column.col) for column in plan.value_columns}
    hierarchy_columns = list(plan.hierarchy_columns)
    for col_idx in range(table_min_col, first_value_col):
        if col_idx in hierarchy_col_indexes or col_idx in value_col_indexes:
            continue
        if not _column_has_data_value(sheet, col_idx, data_min_row, data_max_row):
            continue
        hierarchy_columns.append(
            HierarchyColumn(
                header=_infer_header_for_column(sheet, col_idx, plan),
                col=get_column_letter(col_idx),
            )
        )
    hierarchy_columns.sort(key=lambda item: column_index_from_string(item.col))
    return plan.model_copy(
        update={
            "hierarchy_columns": hierarchy_columns,
            "row_paths": row_paths,
        }
    )


def validate_complex_parse_plan(plan: ComplexTableParsePlan, sheet: Worksheet) -> list[str]:
    warnings: list[str] = []
    _assert_range_inside_sheet(plan.table_range, sheet, "table_range")
    _assert_range_inside_sheet(plan.data_row_range, sheet, "data_row_range")
    for index, range_ref in enumerate(plan.title_ranges):
        _assert_range_inside_sheet(range_ref, sheet, f"title_ranges[{index}]")
    for index, range_ref in enumerate(plan.header_ranges):
        _assert_range_inside_sheet(range_ref, sheet, f"header_ranges[{index}]")

    table_min_col, _, table_max_col, _ = range_boundaries(plan.table_range)
    data_min_col, data_min_row, data_max_col, data_max_row = range_boundaries(plan.data_row_range)
    if data_min_row > data_max_row:
        raise ValueError("data_row_range has no rows")
    if data_min_col < table_min_col or data_max_col > table_max_col:
        warnings.append("data_row_range extends outside table_range columns")

    data_rows = set(range(data_min_row, data_max_row + 1))
    seen_row_paths: set[int] = set()
    for row_idx in plan.ignored_rows:
        if row_idx not in data_rows:
            warnings.append(f"ignored_rows contains row outside data_row_range: {row_idx}")
    for row_path in plan.row_paths:
        if row_path.row not in data_rows:
            raise ValueError(f"row_paths contains row outside data_row_range: {row_path.row}")
        if row_path.row in seen_row_paths:
            warnings.append(f"Duplicate row_path for row: {row_path.row}")
        seen_row_paths.add(row_path.row)
        for segment in row_path.path:
            if segment.ref:
                _assert_cell_ref_inside_sheet(segment.ref, sheet, "row_paths.ref")

    seen_value_cols: set[str] = set()
    for column in [*plan.hierarchy_columns, *plan.value_columns]:
        col_idx = column_index_from_string(column.col)
        if col_idx < table_min_col or col_idx > table_max_col:
            raise ValueError(f"Column {column.col} is outside table_range")
        if col_idx > sheet.max_column:
            raise ValueError(f"Column {column.col} is outside the worksheet")
    for value_column in plan.value_columns:
        if value_column.col in seen_value_cols:
            warnings.append(f"Duplicate value column: {value_column.col}")
        seen_value_cols.add(value_column.col)
    return warnings


def build_complex_plan_result(
    sheet: Worksheet,
    *,
    file_name: str,
    plan: ComplexTableParsePlan,
    raw_plan_output: str,
    validation_warnings: list[str],
) -> PlanBuildResult:
    tree_refs, tree, coverage = build_tree_from_complex_plan(sheet, plan)
    table_min_col, _, table_max_col, _ = range_boundaries(plan.table_range)
    _, data_min_row, _, data_max_row = range_boundaries(plan.data_row_range)
    ignored_rows = set(plan.ignored_rows)
    headers = _headers_from_complex_plan(
        sheet,
        plan=plan,
        min_col=table_min_col,
        max_col=table_max_col,
    )
    records: list[dict[str, Any]] = []
    for row_idx in range(data_min_row, data_max_row + 1):
        if row_idx in ignored_rows:
            continue
        records.append(
            {
                headers[col_idx - table_min_col]: _json_safe(
                    get_effective_cell_value(sheet, row_idx, col_idx)
                )
                for col_idx in range(table_min_col, table_max_col + 1)
            }
        )
    title = _title_from_complex_plan(sheet, plan=plan, fallback=file_name)
    tree_path_text = _flatten_tree_paths(tree, title=title)
    return PlanBuildResult(
        title=title,
        sheet_name=str(sheet.title),
        headers=headers,
        records=records,
        tree=tree,
        tree_with_cell_refs=tree_refs,
        tree_path_text=tree_path_text,
        tree_metric_names=_dedupe_strings([column.header for column in plan.value_columns]),
        row_count=len(records),
        column_count=len(headers),
        parse_plan=plan,
        coverage=coverage,
        validation_warnings=list(validation_warnings),
        raw_plan_output=raw_plan_output,
    )


def build_tree_from_complex_plan(
    sheet: Worksheet,
    plan: ComplexTableParsePlan,
) -> tuple[dict[str, Any], dict[str, Any], CoverageReport]:
    _, data_min_row, _, data_max_row = range_boundaries(plan.data_row_range)
    tree_refs: dict[str, Any] = {}
    missing_cells: list[str] = []
    skipped_rows: list[int] = []
    covered_cells = 0
    last_hierarchy_values: dict[str, str] = {}
    row_path_map = {row_path.row: row_path.path for row_path in plan.row_paths}
    ignored_rows = set(plan.ignored_rows)

    for row_idx in range(data_min_row, data_max_row + 1):
        if row_idx in ignored_rows:
            skipped_rows.append(row_idx)
            continue

        segments = row_path_map.get(row_idx)
        if segments:
            current, row_path_parts = _build_current_from_row_path(tree_refs, sheet, segments)
        else:
            current = tree_refs
            row_path_parts = 0
            previous_hierarchy_value = ""
            for hierarchy_column in plan.hierarchy_columns:
                col_idx = column_index_from_string(hierarchy_column.col)
                raw_value = get_effective_cell_value(sheet, row_idx, col_idx)
                display_value = _display(raw_value)
                if not display_value and plan.hierarchy_fill_down:
                    display_value = last_hierarchy_values.get(hierarchy_column.col, "")
                elif display_value:
                    last_hierarchy_values[hierarchy_column.col] = display_value
                if not display_value:
                    continue
                if (
                    previous_hierarchy_value
                    and display_value == previous_hierarchy_value
                    and _is_merged_child_cell(sheet, row_idx, col_idx)
                ):
                    continue
                semantic_key = f"{hierarchy_column.header} - {display_value}"
                current = current.setdefault(semantic_key, {})
                previous_hierarchy_value = display_value
                row_path_parts += 1

        if row_path_parts == 0:
            skipped_rows.append(row_idx)
            missing_cells.extend(f"{column.col}{row_idx}" for column in plan.value_columns)
            continue

        for value_column in plan.value_columns:
            col_idx = column_index_from_string(value_column.col)
            ref = f"{get_column_letter(col_idx)}{row_idx}"
            if value_column.group:
                current.setdefault(value_column.group, {})[value_column.header] = ref
            else:
                current[value_column.header] = ref
            covered_cells += 1

    data_rows = max(
        0,
        data_max_row
        - data_min_row
        + 1
        - len([row for row in ignored_rows if data_min_row <= row <= data_max_row]),
    )
    coverage = CoverageReport(
        data_rows=data_rows,
        value_columns=len(plan.value_columns),
        expected_cells=data_rows * len(plan.value_columns),
        covered_cells=covered_cells,
        missing_cells=missing_cells,
        skipped_rows=skipped_rows,
    )
    return tree_refs, convert_refs_to_values(tree_refs, sheet), coverage


def convert_refs_to_values(data: dict[str, Any], sheet: Worksheet) -> dict[str, Any]:
    def convert(value: Any) -> Any:
        if isinstance(value, dict):
            return {key: convert(item) for key, item in value.items()}
        if isinstance(value, list):
            return [convert(item) for item in value]
        if isinstance(value, str) and _is_cell_ref(value):
            col_idx, row_idx = _cell_ref_to_indexes(value)
            return _json_safe(get_effective_cell_value(sheet, row_idx, col_idx))
        return value

    return convert(data)


def get_effective_cell_value(sheet: Worksheet, row_idx: int, col_idx: int) -> Any:
    cell = sheet.cell(row_idx, col_idx)
    if cell.value is not None:
        return cell.value
    for merged_range in sheet.merged_cells.ranges:
        if cell.coordinate in merged_range:
            return sheet.cell(merged_range.min_row, merged_range.min_col).value
    return None


def _build_current_from_row_path(
    root: dict[str, Any],
    sheet: Worksheet,
    segments: list[RowPathSegment],
) -> tuple[dict[str, Any], int]:
    current = root
    row_path_parts = 0
    for segment in segments:
        display_value = _display(segment.value)
        if not display_value and segment.ref:
            col_idx, row_idx = _cell_ref_to_indexes(segment.ref)
            display_value = _display(get_effective_cell_value(sheet, row_idx, col_idx))
        if not display_value:
            continue
        semantic_key = f"{segment.header} - {display_value}"
        current = current.setdefault(semantic_key, {})
        row_path_parts += 1
    return current, row_path_parts


def _repair_row_paths(row_paths: list[RowPath], sheet: Worksheet) -> list[RowPath]:
    repaired_paths: list[RowPath] = []
    previous_parent_segments: list[RowPathSegment] = []
    for row_path in sorted(row_paths, key=lambda item: item.row):
        segments = _deduplicate_row_path_segments(row_path.path, sheet)
        qualifier_index = _first_qualifier_index(segments)
        if qualifier_index is not None and previous_parent_segments:
            prefix = [
                segment
                for segment in segments[:qualifier_index]
                if not _is_code_segment(segment) and not _is_qualifier_segment(segment)
            ]
            missing_parent_tail = _missing_parent_tail(prefix, previous_parent_segments, sheet)
            if missing_parent_tail:
                segments = [
                    *segments[:qualifier_index],
                    *missing_parent_tail,
                    *segments[qualifier_index:],
                ]
        repaired_paths.append(row_path.model_copy(update={"path": segments}))
        if qualifier_index is None:
            parent_segments = [
                segment
                for segment in segments
                if not _is_code_segment(segment) and not _is_qualifier_segment(segment)
            ]
            if parent_segments:
                previous_parent_segments = parent_segments
    return repaired_paths


def _deduplicate_row_path_segments(
    segments: list[RowPathSegment],
    sheet: Worksheet,
) -> list[RowPathSegment]:
    deduped: list[RowPathSegment] = []
    for segment in segments:
        if deduped and _same_semantic_segment(deduped[-1], segment, sheet):
            if _is_qualifier_segment(segment) and not _is_qualifier_segment(deduped[-1]):
                deduped[-1] = deduped[-1].model_copy(update={"role": segment.role})
            continue
        deduped.append(segment)
    return deduped


def _first_qualifier_index(segments: list[RowPathSegment]) -> int | None:
    for index, segment in enumerate(segments):
        if _is_qualifier_segment(segment):
            return index
    return None


def _missing_parent_tail(
    prefix: list[RowPathSegment],
    previous_parent_segments: list[RowPathSegment],
    sheet: Worksheet,
) -> list[RowPathSegment]:
    if not prefix:
        return previous_parent_segments
    if len(prefix) > len(previous_parent_segments):
        return []
    for index, segment in enumerate(prefix):
        if not _same_semantic_segment(segment, previous_parent_segments[index], sheet):
            return []
    return previous_parent_segments[len(prefix) :]


def _same_semantic_segment(
    left: RowPathSegment,
    right: RowPathSegment,
    sheet: Worksheet,
) -> bool:
    return left.header == right.header and _segment_display_value(
        left, sheet
    ) == _segment_display_value(right, sheet)


def _segment_display_value(segment: RowPathSegment, sheet: Worksheet) -> str:
    display_value = _display(segment.value)
    if display_value or not segment.ref:
        return display_value
    col_idx, row_idx = _cell_ref_to_indexes(segment.ref)
    return _display(get_effective_cell_value(sheet, row_idx, col_idx))


def _is_code_segment(segment: RowPathSegment) -> bool:
    return (segment.role or "").strip().lower() == "code"


def _is_qualifier_segment(segment: RowPathSegment) -> bool:
    return (segment.role or "").strip().lower() in {
        "qualifier",
        "detail",
        "subitem",
        "sub_item",
        "modifier",
    }


def _column_has_data_value(
    sheet: Worksheet,
    col_idx: int,
    data_min_row: int,
    data_max_row: int,
) -> bool:
    return any(
        _display(get_effective_cell_value(sheet, row_idx, col_idx))
        for row_idx in range(data_min_row, data_max_row + 1)
    )


def _infer_header_for_column(
    sheet: Worksheet,
    col_idx: int,
    plan: ComplexTableParsePlan,
) -> str:
    for range_ref in reversed(plan.header_ranges):
        min_col, min_row, max_col, max_row = range_boundaries(range_ref)
        if not (min_col <= col_idx <= max_col):
            continue
        parts: list[str] = []
        for row_idx in range(min_row, max_row + 1):
            display_value = _display(get_effective_cell_value(sheet, row_idx, col_idx))
            if display_value and display_value not in parts:
                parts.append(display_value)
        if parts:
            return " - ".join(reversed(parts))
    return get_column_letter(col_idx)


def _is_merged_child_cell(sheet: Worksheet, row_idx: int, col_idx: int) -> bool:
    coordinate = sheet.cell(row_idx, col_idx).coordinate
    for merged_range in sheet.merged_cells.ranges:
        if coordinate in merged_range:
            return row_idx != merged_range.min_row or col_idx != merged_range.min_col
    return False


def _build_merged_lookup(sheet: Worksheet) -> dict[tuple[int, int], dict[str, Any]]:
    lookup: dict[tuple[int, int], dict[str, Any]] = {}
    for merged_range in sheet.merged_cells.ranges:
        top_left_value = sheet.cell(merged_range.min_row, merged_range.min_col).value
        for row_idx in range(merged_range.min_row, merged_range.max_row + 1):
            for col_idx in range(merged_range.min_col, merged_range.max_col + 1):
                lookup[(row_idx, col_idx)] = {
                    "range": str(merged_range),
                    "top_left_value": top_left_value,
                    "is_top_left": row_idx == merged_range.min_row
                    and col_idx == merged_range.min_col,
                }
    return lookup


def _assert_range_inside_sheet(
    range_ref: str,
    sheet: Worksheet,
    field_name: str,
) -> None:
    try:
        min_col, min_row, max_col, max_row = range_boundaries(range_ref)
    except Exception as exc:
        raise ValueError(f"{field_name} is not a valid Excel range: {range_ref}") from exc
    if min_col < 1 or min_row < 1 or max_col > sheet.max_column or max_row > sheet.max_row:
        raise ValueError(f"{field_name} is outside the worksheet: {range_ref}")


def _assert_cell_ref_inside_sheet(ref: str, sheet: Worksheet, field_name: str) -> None:
    col_idx, row_idx = _cell_ref_to_indexes(ref)
    if col_idx < 1 or row_idx < 1 or col_idx > sheet.max_column or row_idx > sheet.max_row:
        raise ValueError(f"{field_name} is outside the worksheet: {ref}")


def _cell_ref_to_indexes(ref: str) -> tuple[int, int]:
    match = re.fullmatch(r"\$?([A-Z]+)\$?([0-9]+)", ref.strip().upper())
    if match is None:
        raise ValueError(f"Invalid cell reference: {ref}")
    return column_index_from_string(match.group(1)), int(match.group(2))


def _headers_from_complex_plan(
    sheet: Worksheet,
    *,
    plan: ComplexTableParsePlan,
    min_col: int,
    max_col: int,
) -> list[str]:
    named_columns = {
        column_index_from_string(column.col): column.header
        for column in [*plan.hierarchy_columns, *plan.value_columns]
    }
    headers: list[str] = []
    seen: dict[str, int] = {}
    for col_idx in range(min_col, max_col + 1):
        base = str(named_columns.get(col_idx) or _infer_header_for_column(sheet, col_idx, plan)).strip()
        base = base or get_column_letter(col_idx)
        count = seen.get(base, 0)
        seen[base] = count + 1
        headers.append(base if count == 0 else f"{base}_{count + 1}")
    return headers


def _title_from_complex_plan(
    sheet: Worksheet,
    *,
    plan: ComplexTableParsePlan,
    fallback: str,
) -> str:
    for range_ref in plan.title_ranges:
        min_col, min_row, max_col, max_row = range_boundaries(range_ref)
        values: list[str] = []
        for row_idx in range(min_row, max_row + 1):
            for col_idx in range(min_col, max_col + 1):
                value = _display(get_effective_cell_value(sheet, row_idx, col_idx))
                if value and value not in values:
                    values.append(value)
        if values:
            return " ".join(values)
    header_row = range_boundaries(plan.header_ranges[0])[1] if plan.header_ranges else 1
    table_min_col, _, table_max_col, _ = range_boundaries(plan.table_range)
    return TableParsePlanService._detect_title(
        sheet,
        header_row=header_row,
        min_col=table_min_col,
        max_col=table_max_col,
        fallback=fallback,
    )


def _flatten_tree_paths(tree: dict[str, Any], *, title: str, limit: int = 2000) -> list[str]:
    paths: list[str] = []

    def walk(value: Any, path: list[str]) -> None:
        if len(paths) >= limit:
            return
        if isinstance(value, dict):
            for key, item in value.items():
                walk(item, [*path, str(key)])
            return
        if isinstance(value, list):
            for index, item in enumerate(value, start=1):
                walk(item, [*path, f"第{index}项"])
            return
        prefix = " | ".join([title, *path[:-1]]) if title else " | ".join(path[:-1])
        leaf = f"{path[-1]}: {_display(value)}" if path else _display(value)
        paths.append(" | ".join(part for part in [prefix, leaf] if part))

    walk(tree, [])
    return _dedupe_strings(paths, limit=limit)


def _dedupe_strings(values: list[str], limit: int = 2000) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        normalized = str(value).strip()
        if not normalized or normalized in seen:
            continue
        seen.add(normalized)
        result.append(normalized)
        if len(result) >= limit:
            break
    return result


def _message_to_text(message: BaseMessage | Any) -> str:
    content = getattr(message, "content", message)
    if isinstance(content, str):
        return _extract_after_think(content)
    if isinstance(content, list):
        parts: list[str] = []
        for item in content:
            if isinstance(item, str):
                parts.append(item)
            elif isinstance(item, dict) and isinstance(item.get("text"), str):
                parts.append(item["text"])
        return _extract_after_think("\n".join(parts))
    return _extract_after_think(str(content))


def _extract_after_think(response: str) -> str:
    marker = "</think>"
    return response.split(marker, 1)[1].strip() if marker in response else response.strip()


def _repair_common_json_issues(json_content: str) -> str:
    repaired = json_content.strip()
    repaired = re.sub(r",\s*([}\]])", r"\1", repaired)
    repaired = repaired.replace("“", '"').replace("”", '"').replace("‘", "'").replace("’", "'")
    return repaired


def _is_cell_ref(value: str) -> bool:
    return bool(re.fullmatch(r"\$?[A-Z]+\$?[0-9]+", value.strip().upper()))


table_parse_plan_service = TableParsePlanService()
